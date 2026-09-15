import csv
import openpyxl

wb = openpyxl.load_workbook("Recruit_Monitoring_Tracker.xlsx", data_only=True)
ws = wb["Recruit Monitoring Tracker"]

with open("Recruit_Monitoring_Export.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    # Header
    writer.writerow([ws.cell(row=8, column=c).value for c in range(1, 20)])
    # Rows 9 to 14 (6 encoded recruits)
    for r in range(9, 15):
        writer.writerow([ws.cell(row=r, column=c).value for c in range(1, 20)])

print("Exported Recruit_Monitoring_Export.csv successfully!")
