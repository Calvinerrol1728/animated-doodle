#!/usr/bin/env python3
"""Rebuild the My Agency Game Plan deck (14 slides) with the full game plan filled in.

Recreated from the original Google Slides deck structure:
 1 Business Plan Requirements      8 Recruitment Plan - Sources
 2 My Agency Game Plan (title)     9 Recruitment Plan - Recruitment Activities
 3 VISION                         10 Break slide
 4 MISSION                        11 Activation Plan
 5 CORE VALUES (CARE)             12 Structured Work Week
 6 GOAL SHEET                     13 Thank You
 7 Recruitment Plan - Profile     14 My Commitment (original: empty "TEXT" slide)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
NAVY      = RGBColor(0x0A, 0x2A, 0x5E)
NAVY_DEEP = RGBColor(0x07, 0x1E, 0x42)
RED       = RGBColor(0xD6, 0x00, 0x1C)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x21, 0x25, 0x2B)
GRAY      = RGBColor(0x5A, 0x62, 0x70)
GRAY_LT   = RGBColor(0xF2, 0xF4, 0xF8)
LINE_LT   = RGBColor(0xD8, 0xDE, 0xE8)
MIST      = RGBColor(0xB9, 0xC4, 0xD6)   # light blue-gray on navy
SUN_TINT  = RGBColor(0xED, 0xF1, 0xF7)
FONT      = "Arial"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def slide_new(bg=None):
    s = prs.slides.add_slide(BLANK)
    if bg is not None:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
    return s

def tb(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def para(tf, text=None, first=True, align=PP_ALIGN.LEFT, space_after=None,
         line_spacing=None):
    p0 = tf.paragraphs[0]
    # use the first paragraph only while it is still empty; otherwise append
    p = p0 if (first and not p0.runs) else tf.add_paragraph()
    p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    return p

def run(p, text, size=12, bold=False, color=INK, italic=False, font=FONT):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    return r

def rect(slide, x, y, w, h, fill, line=None, rounded=False, radius=0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp

def header(slide, kicker, title):
    t = tb(slide, 0.57, 0.30, 12.2, 0.32)
    run(para(t), kicker.upper(), 13, True, RED)
    t = tb(slide, 0.55, 0.60, 12.2, 0.64)
    run(para(t), title, 29, True, NAVY)
    rect(slide, 0.57, 1.30, 1.5, 0.055, RED)

def footer(slide):
    t = tb(slide, 0.55, 7.13, 6.5, 0.26)
    run(para(t), "My Agency Game Plan — Presentation 2026", 9, False, GRAY)
    t = tb(slide, 7.0, 7.13, 5.78, 0.26)
    run(para(t, align=PP_ALIGN.RIGHT), "Agency Builder Onboarding Program", 9, False, GRAY)

def bullet(tf, text, first=False, size=12.5, color=INK, space=7, bold_head=None,
           line_spacing=1.08):
    p = para(tf, first=first, space_after=space, line_spacing=line_spacing)
    run(p, "•  ", size, True, RED)
    if bold_head:
        run(p, bold_head, size, True, color)
    run(p, text, size, False, color)
    return p

def set_cell(cell, text, size=11, bold=False, color=INK, fill=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color

def make_table(slide, x, y, w, rows, cols, col_widths, header_h, row_h):
    gf = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w),
                                Inches(header_h + row_h * (rows - 1)))
    t = gf.table
    t.first_row = False
    t.horz_banding = False
    for i, cw in enumerate(col_widths):
        t.columns[i].width = Inches(cw)
    t.rows[0].height = Inches(header_h)
    for i in range(1, rows):
        t.rows[i].height = Inches(row_h)
    return t

# ================================================================ SLIDE 1
s = slide_new()
header(s, "Agency Builder Onboarding Program", "Business Plan Requirements")
footer(s)

t = tb(s, 0.55, 1.52, 7.55, 5.5)
run(para(t, space_after=8), "Your Business Plan", 16, True, NAVY)
bullet(t, "Fill out your Business Plan Template. (Instructions will be given during training)", first=True, size=12.5, space=8)
bullet(t, "During the 3-day training program, you will strengthen and improve this plan using the tools, frameworks, and discussions from each session. Your final Business Plan will also be used as part of your Day 3 validation activity.", size=12.5, space=8)
bullet(t, "The purpose of this activity is to help you:", size=12.5, space=3)
for sub in ["Clarify your goals as a new Unit Head",
            "Start thinking about your recruitment strategy",
            "Apply the learnings immediately during the program",
            "Build a practical plan that you can execute after training"]:
    p = para(t, space_after=3, line_spacing=1.05)
    run(p, "      –  ", 12, False, GRAY)
    run(p, sub, 12, False, GRAY)
s2_ = para(t, space_after=8)
run(s2_, "", 6)
bullet(t, "At the end of the program, you will present your enhanced Business Plan and explain how you will build and activate your team.", size=12.5, space=10)
run(para(t, space_after=4), "Note", 13, True, NAVY)
for note in ["Your Business Plan does not need to be perfect — it is a working document.",
             "Think of this as your first draft. Throughout the program, you will refine and strengthen it using the tools, insights, and best practices from each session.",
             "By Day 3, you will present your enhanced Business Plan at the Business Plan Showcase."]:
    p = para(t, space_after=4, line_spacing=1.05)
    run(p, "–  ", 11, False, GRAY)
    run(p, note, 11, False, GRAY, italic=True)

card = rect(s, 8.45, 1.52, 4.33, 5.42, GRAY_LT, line=NAVY, rounded=True, radius=0.045)
t = tb(s, 8.75, 1.82, 3.75, 4.9)
run(para(t, space_after=6), "Your Checklist", 16, True, NAVY)
run(para(t, space_after=10), "I have:", 12, False, INK)
for item in ["Completed my Vision", "Completed my Mission", "Completed my Core Values",
             "Completed my Goal Sheet", "Completed my Recruitment Sources",
             "Completed my Recruitment Activities", "Reviewed my Business Plan"]:
    p = para(t, space_after=9, line_spacing=1.0)
    run(p, "☐  ", 13, True, NAVY)
    run(p, item, 12.5, False, INK)

# ================================================================ SLIDE 2
s = slide_new(NAVY_DEEP)
rect(s, 0.9, 2.00, 1.7, 0.07, RED)
t = tb(s, 0.87, 2.18, 11.5, 1.15)
run(para(t), "My Agency", 54, True, WHITE)
t = tb(s, 0.87, 3.28, 11.5, 1.15)
run(para(t), "Game Plan", 54, True, WHITE)
t = tb(s, 0.9, 4.62, 9, 0.45)
run(para(t), "PRESENTATION 2026", 15, True, MIST)
t = tb(s, 0.9, 6.52, 10, 0.4)
run(para(t), "Prepared by:  [Your Name]   •   Agency Builder, Unit Head Candidate", 13, False, MIST)

# ================================================================ SLIDE 3 — VISION
s = slide_new(NAVY)
rect(s, 0.9, 1.30, 1.4, 0.07, RED)
t = tb(s, 0.88, 1.48, 5, 0.5)
run(para(t), "VISION", 17, True, WHITE)
t = tb(s, 0.88, 2.20, 11.4, 3.7)
run(para(t, line_spacing=1.18),
    "By 2030, we will have achieved sustainable growth and client success at "
    "scale — becoming the most distinctive and admired team in our industry "
    "for our innovation, quality, and people-centric culture.", 30, True, WHITE)
t = tb(s, 0.9, 6.55, 11, 0.4)
run(para(t), "Our North Star — every goal in this plan ladders up to this vision.", 12, False, MIST, italic=True)

# ================================================================ SLIDE 4 — MISSION
s = slide_new(NAVY)
rect(s, 0.9, 1.30, 1.4, 0.07, RED)
t = tb(s, 0.88, 1.48, 5, 0.5)
run(para(t), "MISSION", 17, True, WHITE)
t = tb(s, 0.88, 2.20, 11.4, 3.7)
run(para(t, line_spacing=1.18),
    "We exist to develop successful financial advisors and future leaders who "
    "transform lives. We add value by empowering families to live with "
    "confidence, purpose, and financial freedom — through proactive and "
    "disciplined financial planning.", 30, True, WHITE)
t = tb(s, 0.9, 6.55, 11, 0.4)
run(para(t), "Every activity in this Game Plan exists to deliver this Mission.", 12, False, MIST, italic=True)

# ================================================================ SLIDE 5 — CORE VALUES
s = slide_new()
header(s, "The Way We Work", "CORE VALUES")
footer(s)
values = [
    ("C", "Client Obsessed",
     "Clients first, always — we listen, advise, and deliver on our promises."),
    ("A", "Accountability and Integrity",
     "We own our numbers, keep our word, and do the right thing even when no one is watching."),
    ("R", "Resilience and Courage",
     "We bounce back from every \u2018no\u2019 and show up braver the next day."),
    ("E", "Empowerment and One Team",
     "We grow leaders, lift each other up, and win as one team."),
]
y = 1.62
for i, (letter, name, desc) in enumerate(values):
    t = tb(s, 0.60, y, 1.0, 1.05)
    run(para(t, align=PP_ALIGN.CENTER), letter, 46, True, RED)
    t = tb(s, 1.80, y + 0.10, 6.5, 0.5)
    run(para(t), name, 21, True, NAVY)
    t = tb(s, 1.80, y + 0.56, 10.6, 0.45)
    run(para(t), desc, 12.5, False, GRAY)
    if i < 3:
        rect(s, 1.80, y + 1.13, 10.7, 0.014, LINE_LT)
    y += 1.34

# ================================================================ SLIDE 6 — GOAL SHEET
s = slide_new()
header(s, "GOAL SHEET", "My Goals — From 2030 Vision to 90-Day Sprint")
footer(s)
cards = [
    ("2030 NORTH STAR", NAVY, [
        "The most distinctive and admired team in the industry",
        "A self-sustaining agency of 25+ productive advisors",
        "1,000+ families across Cebu & Central Visayas living with financial confidence",
        "Leaders we developed, now leading their own teams",
    ]),
    ("12-MONTH TARGETS", RED, [
        "12 recruits licensed — 1 per month",
        "8 active producers by December",
        "120+ team cases in the year",
        "100 new families served",
        "Personal production: 2 cases per month (24 for the year)",
    ]),
    ("90-DAY SPRINT", NAVY, [
        "Build my 100-name candidate list",
        "30 recruitment conversations → 3 recruits",
        "3 recruits contracted and licensed",
        "6 personal cases closed",
        "Weekly scorecard live from Week 1",
    ]),
]
x = 0.55
for title, band, items in cards:
    rect(s, x, 1.60, 4.0, 4.90, WHITE, line=LINE_LT, rounded=True, radius=0.03)
    rect(s, x, 1.60, 4.0, 0.14, band)
    t = tb(s, x + 0.26, 1.90, 3.5, 0.4)
    run(para(t), title, 16, True, NAVY)
    t = tb(s, x + 0.26, 2.38, 3.5, 3.95)
    for i, item in enumerate(items):
        bullet(t, item, first=(i == 0), size=12, color=INK, space=8, line_spacing=1.05)
    x += 4.11
t = tb(s, 0.55, 6.68, 12.2, 0.35)
run(para(t), "Reviewed every Friday  •  Adjusted monthly  •  Shared with my team leader for accountability", 11.5, False, GRAY, italic=True)

# ================================================================ SLIDE 7 — PROFILE OF RECRUIT
s = slide_new()
header(s, "Recruitment Plan", "Profile of Recruit")
footer(s)
t = tb(s, 0.55, 1.50, 12.2, 0.4)
run(para(t), "What are your ideal candidate's qualifications, skills, experience, education, and personal attributes?", 12.5, False, GRAY, italic=True)

rect(s, 0.55, 2.00, 7.35, 4.90, GRAY_LT, rounded=True, radius=0.03)
t = tb(s, 0.85, 2.24, 6.8, 0.4)
run(para(t), "Ideal Candidate Profile", 15, True, NAVY)
t = tb(s, 0.85, 2.76, 6.8, 4.0)
profiles = [
    "College graduate, any course (ages 23–45)",
    "1–3 years in sales, service, teaching, BPO, banking, or real estate",
    "Entrepreneurial mindset — hungry for unlimited income growth",
    "Strong communication and relationship-building skills",
    "Coachable, disciplined, and self-motivated",
    "Based in Cebu / Central Visayas and open to field work",
    "Newly licensed — or willing to take the licensing exam",
]
for i, item in enumerate(profiles):
    bullet(t, item, first=(i == 0), size=12.5, space=8, line_spacing=1.05)

rect(s, 8.15, 2.00, 4.63, 4.90, NAVY, rounded=True, radius=0.03)
t = tb(s, 8.45, 2.24, 4.0, 0.4)
run(para(t), "CARE Fit Checklist", 15, True, WHITE)
fit = [
    ("C", "Client Obsessed", "Genuinely enjoys helping families win"),
    ("A", "Accountable", "Honest with money, records, and time"),
    ("R", "Resilient", "Handles rejection without losing drive"),
    ("E", "One Team", "Team player who celebrates others' wins"),
]
yy = 2.90
for letter, label, desc in fit:
    t = tb(s, 8.45, yy, 4.1, 0.9)
    p = para(t, space_after=2)
    run(p, f"{letter} — ", 13.5, True, RGBColor(0xFF, 0x6B, 0x7D))
    run(p, label, 13.5, True, WHITE)
    p = para(t)
    run(p, desc, 11, False, MIST)
    yy += 1.0

# ================================================================ SLIDE 8 — SOURCES
s = slide_new()
header(s, "Recruitment Plan", "Sources")
footer(s)
t = tb(s, 0.55, 1.50, 12.2, 0.4)
run(para(t), "Identify your sources — e.g., social media platforms, family, friends, referrals, etc.", 12.5, False, GRAY, italic=True)
sources = [
    ("Warm Market", "Family, relatives, friends, classmates, and neighbors — the first 50 names on my list."),
    ("Client & COI Referrals", "Ask for 2 names at every client meeting — doctors, bankers, business owners."),
    ("Social Media", "Facebook, Instagram, LinkedIn, TikTok — 3 career posts and daily stories."),
    ("Community & Orgs", "Church, alumni association, barangay groups, and business associations."),
    ("Past Career Network", "Ex-colleagues from BPO, sales, teaching, and banking who want more."),
    ("AXA Discovery Day", "Invite 2 guests per run; convert every attendee to a coffee chat."),
]
xs = [0.55, 4.66, 8.77]
ys = [2.05, 4.35]
for i, (title, desc) in enumerate(sources):
    cx, cy = xs[i % 3], ys[i // 3]
    rect(s, cx, cy, 4.0, 2.1, WHITE, line=LINE_LT, rounded=True, radius=0.05)
    rect(s, cx, cy, 0.09, 2.1, RED)
    t = tb(s, cx + 0.3, cy + 0.22, 3.5, 0.4)
    run(para(t), title, 14, True, NAVY)
    t = tb(s, cx + 0.3, cy + 0.68, 3.5, 1.3)
    run(para(t, line_spacing=1.12), desc, 11, False, INK)

# ================================================================ SLIDE 9 — RECRUITMENT ACTIVITIES
s = slide_new()
header(s, "Recruitment Plan", "Recruitment Activities")
footer(s)
t = tb(s, 0.55, 1.48, 12.2, 0.4)
run(para(t), "The specific actions I will regularly do to attract, engage, and meet potential candidates:", 12.5, False, GRAY, italic=True)
rows9 = [
    ("Set recruitment appointments", "30 mins a day [Mon to Sat]", "10 per day"),
    ("Coffee chats with warm market", "5 invitations per day", "25 invites per week"),
    ("Career posts & stories", "3 posts per week (FB / IG / LinkedIn)", "15 quality leads per month"),
    ("Referral asks", "At every client meeting", "2 names per meeting"),
    ("Discovery Day guests", "2 guests per run (2× per month)", "4 guests per month"),
    ("Candidate follow-ups", "Within 48 hours of first meeting", "100% followed up"),
    ("Campus / community career talks", "1 talk per month", "20 attendees per month"),
]
tbl = make_table(s, 0.55, 2.00, 12.23, 8, 3, [4.63, 4.10, 3.50], 0.48, 0.56)
for c, htxt in enumerate(["Activity", "Frequency", "Measure of Success"]):
    set_cell(tbl.cell(0, c), htxt, 12.5, True, WHITE, NAVY)
for r, (a, b, c) in enumerate(rows9, start=1):
    fill = WHITE if r % 2 else GRAY_LT
    set_cell(tbl.cell(r, 0), a, 11.5, True, INK, fill)
    set_cell(tbl.cell(r, 1), b, 11.5, False, INK, fill)
    set_cell(tbl.cell(r, 2), c, 11.5, False, NAVY, fill)

# ================================================================ SLIDE 10 — BREAK
s = slide_new(NAVY_DEEP)
rect(s, 6.17, 2.42, 1.0, 0.07, RED)
t = tb(s, 0.5, 2.75, 12.33, 1.3)
run(para(t, align=PP_ALIGN.CENTER), "BREAK", 66, True, WHITE)
t = tb(s, 0.5, 4.25, 12.33, 0.5)
run(para(t, align=PP_ALIGN.CENTER), "Back in 15 minutes.", 16, False, MIST)

# ================================================================ SLIDE 11 — ACTIVATION PLAN
s = slide_new()
header(s, "ACTIVATION PLAN", "Activation Plan")
footer(s)
t = tb(s, 0.55, 1.48, 12.2, 0.4)
run(para(t), "How I will train, coach, and retain every recruit — especially in their first 90 days:", 12.5, False, GRAY, italic=True)
rows11 = [
    ("New-recruit onboarding + 100-name list", "Week 1 after contracting", "100-name list submitted"),
    ("Licensing & exam study support", "1 hour daily (first 60 days)", "Licensed within 60 days"),
    ("Joint field work with recruit", "2× per week", "8 joint meetings per month"),
    ("One-on-one coaching", "30 minutes weekly", "Weekly scorecard reviewed"),
    ("First-case milestone", "Within 45 days of licensing", "3 cases in first 90 days"),
    ("Team sales meeting & product training", "Every Tuesday", "90% attendance"),
    ("Monthly CARE check-in & recognition", "Once a month", "80% of recruits active at Month 12"),
]
tbl = make_table(s, 0.55, 2.00, 12.23, 8, 3, [4.63, 4.10, 3.50], 0.48, 0.56)
for c, htxt in enumerate(["Activity", "Frequency", "Measure of Success"]):
    set_cell(tbl.cell(0, c), htxt, 12.5, True, WHITE, NAVY)
for r, (a, b, c) in enumerate(rows11, start=1):
    fill = WHITE if r % 2 else GRAY_LT
    set_cell(tbl.cell(r, 0), a, 11.5, True, INK, fill)
    set_cell(tbl.cell(r, 1), b, 11.5, False, INK, fill)
    set_cell(tbl.cell(r, 2), c, 11.5, False, NAVY, fill)

# ================================================================ SLIDE 12 — STRUCTURED WORK WEEK
s = slide_new()
t = tb(s, 0.57, 0.26, 12.2, 0.30)
run(para(t), "DISCIPLINE", 13, True, RED)
t = tb(s, 0.55, 0.54, 12.2, 0.56)
run(para(t), "STRUCTURED WORK WEEK", 27, True, NAVY)
rect(s, 0.57, 1.14, 1.5, 0.05, RED)

days = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY"]
grid = {
    "MONDAY":    ["Team huddle & plan", "Prospecting — 10 calls", "Client meetings", "Client meetings",
                  "Lunch & recharge", "Prospecting — 10 calls", "Joint field work", "Follow-ups & referrals",
                  "Recruit interviews", "Study 30 min + network"],
    "TUESDAY":   ["Team huddle & plan", "Prospecting — 10 calls", "Recruit interviews", "Joint field work",
                  "Lunch & recharge", "Client meetings", "Client meetings", "Referral asks",
                  "Candidate follow-ups", "Study 30 min + network"],
    "WEDNESDAY": ["Team huddle & plan", "Prospecting — 10 calls", "Client meetings", "Follow-ups & referrals",
                  "Lunch & recharge", "Recruit interviews", "Joint field work", "Follow-ups & referrals",
                  "Recruit interviews", "Study 30 min + network"],
    "THURSDAY":  ["Team huddle & plan", "Prospecting — 10 calls", "Recruit interviews", "Follow-ups & referrals",
                  "Lunch & recharge", "Client meetings", "Client meetings", "Referral asks",
                  "Candidate follow-ups", "Study 30 min + network"],
    "FRIDAY":    ["Team huddle & plan", "Prospecting — 10 calls", "Client meetings", "Client meetings",
                  "Lunch & recharge", "New-recruit coaching", "New-recruit coaching", "Admin & CRM update",
                  "Weekly review & scorecard", "Family time"],
    "SATURDAY":  ["Team training", "Team training", "Discovery Day (2×/mo)", "Discovery Day",
                  "Lunch & recharge", "Personal client meetings", "Personal client meetings",
                  "Case service & papering", "Personal study (CPD)", "Family time"],
    "SUNDAY":    ["Rest", "Mass & gratitude", "Family time", "Family time", "Family lunch",
                  "Rest", "Friends & recharge", "Rest", "Plan next week — 30 min", "Early night"],
}
times = ["8:00-9:00", "9:00-10:00", "10:00-11:00", "11:00-12:00", "12:00-1:00",
         "1:00-2:00", "2:00-3:00", "3:00-4:00", "4:00-5:00", "5:00-6:00"]

tbl = make_table(s, 0.35, 1.36, 12.63, 12, 8, [1.15, 1.64, 1.64, 1.64, 1.64, 1.64, 1.64, 1.64],
                 0.40, 0.475)
set_cell(tbl.cell(0, 0), "TIME", 9.5, True, WHITE, NAVY, align=PP_ALIGN.CENTER)
for c, d in enumerate(days, start=1):
    set_cell(tbl.cell(0, c), d, 9.5, True, WHITE, NAVY, align=PP_ALIGN.CENTER)
set_cell(tbl.cell(1, 0), "DATE", 9, True, GRAY, GRAY_LT, align=PP_ALIGN.CENTER)
for c in range(1, 8):
    set_cell(tbl.cell(1, c), "", 9, False, INK, WHITE)
for r, tm in enumerate(times, start=2):
    set_cell(tbl.cell(r, 0), tm, 8.5, True, NAVY, GRAY_LT, align=PP_ALIGN.CENTER)
    for c, d in enumerate(days, start=1):
        fill = SUN_TINT if d == "SUNDAY" else WHITE
        set_cell(tbl.cell(r, c), grid[d][r - 2], 8, False, INK, fill, align=PP_ALIGN.CENTER)

t = tb(s, 0.55, 7.08, 12.2, 0.32)
run(para(t), "This calendar IS the plan — discipline beats motivation. Sunday is protected for family and faith.", 10.5, False, GRAY, italic=True)

# ================================================================ SLIDE 13 — THANK YOU
s = slide_new(NAVY_DEEP)
rect(s, 0.9, 2.30, 1.4, 0.07, RED)
t = tb(s, 0.87, 2.52, 11.5, 1.2)
run(para(t), "Thank You.", 60, True, WHITE)
t = tb(s, 0.9, 3.95, 10, 0.45)
run(para(t), "Prepared by:  [Your Name] — Agency Builder", 15, False, MIST)
t = tb(s, 0.9, 6.55, 11, 0.4)
run(para(t), "My Agency Game Plan  •  Presentation 2026", 12, False, MIST)

# ================================================================ SLIDE 14 — MY COMMITMENT
s = slide_new()
header(s, "My Commitment", "This Is My Game Plan.")
footer(s)
t = tb(s, 0.55, 1.75, 12.2, 1.7)
run(para(t, line_spacing=1.3),
    "I will work this plan with discipline and a servant's heart — leading with "
    "CARE in every client conversation, every recruit I develop, and every family "
    "I serve. I will review my numbers weekly, adjust monthly, and never lower "
    "the standard — only the deadline.", 15.5, False, INK)
t = tb(s, 0.55, 3.85, 8, 0.45)
run(para(t), "[Your Name] — Agency Builder", 14, True, NAVY)
rect(s, 0.55, 4.75, 4.7, 0.018, NAVY)
t = tb(s, 0.55, 4.85, 4.7, 0.35)
run(para(t), "Signature over printed name", 10.5, False, GRAY)
rect(s, 5.85, 4.75, 3.2, 0.018, NAVY)
t = tb(s, 5.85, 4.85, 3.2, 0.35)
run(para(t), "Date", 10.5, False, GRAY)

OUT = "/home/user/animated-doodle/My_Agency_Game_Plan_2026.pptx"
prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
