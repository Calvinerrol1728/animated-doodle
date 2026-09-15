import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter
import datetime

def create_recruit_workbook(filename="Recruit_Monitoring_Tracker.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove initial default sheet

    # Typography & Fonts
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="334155")
    font_sec_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_col_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_data_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_muted = Font(name="Calibri", size=9, color="64748B")

    # Borders
    border_light = Side(border_style="thin", color="CBD5E1")
    border_dark = Side(border_style="medium", color="64748B")
    border_double = Side(border_style="double", color="1E3A8A")
    
    cell_border = Border(left=border_light, right=border_light, top=border_light, bottom=border_light)
    header_border = Border(left=border_light, right=border_light, top=border_light, bottom=border_dark)
    total_border = Border(left=border_light, right=border_light, top=border_light, bottom=border_double)

    # -------------------------------------------------------------
    # 1. SHEET 1: Recruit Monitoring Tracker
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Recruit Monitoring Tracker")
    ws1.views.sheetView[0].showGridLines = True
    ws1.freeze_panes = "C9" # Freezes top 8 rows and columns A-B (Recruit ID & Name)

    # Row 1: Title Banner
    ws1.merge_cells("A1:S1")
    ws1["A1"] = "RECRUIT MONITORING TRACKER"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws1.row_dimensions[1].height = 36

    # Row 2: Subtitle
    ws1.merge_cells("A2:S2")
    ws1["A2"] = "Agent Recruitment Pipeline | Candidate Profiling, ADD (Discovery Day) & LMS (Learning Modules) Tracker"
    ws1["A2"].font = font_subtitle
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws1.row_dimensions[2].height = 20

    # Row 3: Spacing
    ws1.row_dimensions[3].height = 8

    # Helper function to style KPI card blocks
    def style_card(start_col, end_col, bg_color, border_color, text_color, label, val_formula, is_pct, sub_formula_or_text):
        c_side = Side(border_style="thin", color=border_color)
        c_border = Border(left=c_side, right=c_side, top=c_side, bottom=c_side)
        c_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        
        ws1.merge_cells(f"{start_col}4:{end_col}4")
        ws1.merge_cells(f"{start_col}5:{end_col}5")
        ws1.merge_cells(f"{start_col}6:{end_col}6")
        
        top_cell = ws1[f"{start_col}4"]
        top_cell.value = label
        top_cell.font = Font(name="Calibri", size=9, bold=True, color=text_color)
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        mid_cell = ws1[f"{start_col}5"]
        mid_cell.value = val_formula
        mid_cell.font = Font(name="Calibri", size=20, bold=True, color=text_color)
        mid_cell.alignment = Alignment(horizontal="center", vertical="center")
        if is_pct:
            mid_cell.number_format = "0.0%"
        else:
            mid_cell.number_format = "#,##0"

        bot_cell = ws1[f"{start_col}6"]
        bot_cell.value = sub_formula_or_text
        bot_cell.font = Font(name="Calibri", size=8.5, bold=False, color=text_color)
        bot_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        from openpyxl.utils.cell import range_boundaries
        for r in range(4, 7):
            min_col, _, max_col, _ = range_boundaries(f"{start_col}{r}:{end_col}{r}")
            for c in range(min_col, max_col + 1):
                cell = ws1.cell(row=r, column=c)
                cell.fill = c_fill
                cell.border = c_border

    # 5 Executive KPI Summary Cards
    style_card("B", "D", "EBF4FF", "93C5FD", "1E40AF", "TOTAL RECRUITS TRACKED", "=COUNTA(B9:B65)", False, "Active candidates in pipeline")
    style_card("F", "H", "E0F2FE", "7DD3FC", "0369A1", "ATTENDED ADD", '=COUNTIF(H9:H65, "Yes")', False, '="Attendance Rate: " & TEXT(IFERROR(COUNTIF(H9:H65,"Yes")/COUNTA(B9:B65),0),"0.0%")')
    style_card("J", "L", "FEF3C7", "FCD34D", "B45309", "LMS IN PROGRESS", '=COUNTIF(L9:L65, "In Progress")', False, "Currently taking online modules")
    style_card("N", "P", "DCFCE7", "86EFAC", "15803D", "COMPLETED LMS", '=COUNTIF(L9:L65, "Yes")', False, '="Completion Rate: " & TEXT(IFERROR(COUNTIF(L9:L65,"Yes")/COUNTIF(H9:H65,"Yes"),0),"0.0%")')
    style_card("R", "S", "F3E8FF", "D8B4FE", "7E22CE", "FUNNEL CONVERSION", '=IFERROR(COUNTIF(L9:L65,"Yes")/COUNTA(B9:B65),0)', True, "Total Recruits -> LMS Certified")

    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 28
    ws1.row_dimensions[6].height = 18

    # Row 7: Group/Phase Super-Headers
    def set_sec_header(range_str, text, fill_hex):
        ws1.merge_cells(range_str)
        from openpyxl.utils.cell import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(range_str)
        fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
        for c in range(min_col, max_col + 1):
            cell = ws1.cell(row=min_row, column=c)
            cell.font = font_sec_hdr
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border
        ws1[f"{get_column_letter(min_col)}{min_row}"] = text

    set_sec_header("A7:F7", "CANDIDATE INFORMATION", "1B365D")
    set_sec_header("G7:J7", "STAGE 1: ADD (AGENCY DISCOVERY DAY)", "2563EB")
    set_sec_header("K7:O7", "STAGE 2: LMS (LEARNING MANAGEMENT SYSTEM)", "0D9488")
    set_sec_header("P7:S7", "PIPELINE STATUS & NEXT ACTIONS", "334155")
    ws1.row_dimensions[7].height = 22

    # Row 8: Table Column Headers
    headers = [
        ("A", "Recruit ID", "center", "244B7A"),
        ("B", "Complete Name", "left", "244B7A"),
        ("C", "Contact Number", "center", "244B7A"),
        ("D", "Email Address", "left", "244B7A"),
        ("E", "Sourced By", "left", "244B7A"),
        ("F", "Date Sourced", "center", "244B7A"),
        ("G", "ADD Scheduled Date", "center", "3B82F6"),
        ("H", "Attended ADD?", "center", "3B82F6"),
        ("I", "ADD Attendance Date", "center", "3B82F6"),
        ("J", "ADD Remarks / Feedback", "left", "3B82F6"),
        ("K", "LMS Access Date", "center", "14B8A6"),
        ("L", "Completed LMS?", "center", "14B8A6"),
        ("M", "LMS Completion Date", "center", "14B8A6"),
        ("N", "LMS Progress %", "center", "14B8A6"),
        ("O", "LMS Score / Certificate", "left", "14B8A6"),
        ("P", "Current Pipeline Stage", "center", "475569"),
        ("Q", "Next Action Required", "left", "475569"),
        ("R", "Target Follow-up Date", "center", "475569"),
        ("S", "Recruiter Notes", "left", "475569"),
    ]

    for col_letter, title, align, fill_hex in headers:
        cell = ws1[f"{col_letter}8"]
        cell.value = title
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
        cell.border = header_border

    ws1.row_dimensions[8].height = 30

    # User's Encoded Recruits:
    # 1. Maybe Grandes (Attended ADD Sept 14, setting up LMS today)
    # 2. Michael Bagiuos
    # 3. Alona Barilea
    # 4. Lorenz Garrucha
    # 5. Princess Ambos
    # 6. Jezreel Kate Montano
    recruits_list = [
        (
            "REC-001",
            "Grandes, Maybe",
            "", "", "",
            datetime.date(2026, 9, 14),
            datetime.date(2026, 9, 14),
            "Yes",
            datetime.date(2026, 9, 14),
            "Attended ADD on Sept 14; confirmed interest. Preparing for LMS setup.",
            datetime.date(2026, 9, 15),
            "Not Started",
            None,
            0.0,
            "Pending account activation",
            "Set up LMS account & send module access link today",
            datetime.date(2026, 9, 15),
            "Attended ADD yesterday Sept 14, 2026. Will set up her LMS today."
        ),
        (
            "REC-002",
            "Bagiuos, Michael",
            "", "", "",
            datetime.date(2026, 9, 15),
            None,
            "Pending",
            None,
            "Newly added recruit; pending ADD orientation schedule.",
            None,
            "Not Started",
            None,
            0.0,
            "",
            "Confirm ADD schedule & send orientation invite",
            datetime.date(2026, 9, 16),
            "Recruit added to monitoring pipeline."
        ),
        (
            "REC-003",
            "Barilea, Alona",
            "", "", "",
            datetime.date(2026, 9, 15),
            None,
            "Pending",
            None,
            "Newly added recruit; pending ADD orientation schedule.",
            None,
            "Not Started",
            None,
            0.0,
            "",
            "Confirm ADD schedule & send orientation invite",
            datetime.date(2026, 9, 16),
            "Recruit added to monitoring pipeline."
        ),
        (
            "REC-004",
            "Garrucha, Lorenz",
            "", "", "",
            datetime.date(2026, 9, 15),
            None,
            "Pending",
            None,
            "Newly added recruit; pending ADD orientation schedule.",
            None,
            "Not Started",
            None,
            0.0,
            "",
            "Confirm ADD schedule & send orientation invite",
            datetime.date(2026, 9, 16),
            "Recruit added to monitoring pipeline."
        ),
        (
            "REC-005",
            "Ambos, Princess",
            "", "", "",
            datetime.date(2026, 9, 15),
            None,
            "Pending",
            None,
            "Newly added recruit; pending ADD orientation schedule.",
            None,
            "Not Started",
            None,
            0.0,
            "",
            "Confirm ADD schedule & send orientation invite",
            datetime.date(2026, 9, 16),
            "Recruit added to monitoring pipeline."
        ),
        (
            "REC-006",
            "Montano, Jezreel Kate",
            "", "", "",
            datetime.date(2026, 9, 15),
            None,
            "Pending",
            None,
            "Newly added recruit; pending ADD orientation schedule.",
            None,
            "Not Started",
            None,
            0.0,
            "",
            "Confirm ADD schedule & send orientation invite",
            datetime.date(2026, 9, 16),
            "Recruit added to monitoring pipeline."
        ),
    ]

    # Populate Rows 9 to 65
    for r in range(9, 66):
        is_encoded = (r - 9) < len(recruits_list)
        bg_hex = "FFFFFF" if r % 2 == 1 else "F8FAFC"
        row_fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
        ws1.row_dimensions[r].height = 22

        if is_encoded:
            data = recruits_list[r - 9]
            ws1[f"A{r}"] = data[0]
            ws1[f"B{r}"] = data[1]
            ws1[f"C{r}"] = data[2]
            ws1[f"D{r}"] = data[3]
            ws1[f"E{r}"] = data[4]
            ws1[f"F{r}"] = data[5]
            ws1[f"G{r}"] = data[6]
            ws1[f"H{r}"] = data[7]
            ws1[f"I{r}"] = data[8]
            ws1[f"J{r}"] = data[9]
            ws1[f"K{r}"] = data[10]
            ws1[f"L{r}"] = data[11]
            ws1[f"M{r}"] = data[12]
            ws1[f"N{r}"] = data[13]
            ws1[f"O{r}"] = data[14]
            ws1[f"P{r}"] = f'=IF(B{r}="","",IF(L{r}="Yes","LMS Completed (Ready for Exam)",IF(L{r}="In Progress","LMS In Progress",IF(H{r}="Yes","ADD Attended - Awaiting LMS",IF(H{r}="Rescheduled","ADD Rescheduled",IF(H{r}="Pending","ADD Scheduled",IF(H{r}="No","ADD Missed - Follow Up","Initial Prospect")))))))'
            ws1[f"Q{r}"] = data[15]
            ws1[f"R{r}"] = data[16]
            ws1[f"S{r}"] = data[17]
        else:
            # Clean pre-formatted blank template rows
            ws1[f"A{r}"] = f'=IF(B{r}="","","REC-"&TEXT({r}-8,"000"))'
            ws1[f"B{r}"] = ""
            ws1[f"C{r}"] = ""
            ws1[f"D{r}"] = ""
            ws1[f"E{r}"] = ""
            ws1[f"F{r}"] = None
            ws1[f"G{r}"] = None
            ws1[f"H{r}"] = ""
            ws1[f"I{r}"] = None
            ws1[f"J{r}"] = ""
            ws1[f"K{r}"] = None
            ws1[f"L{r}"] = ""
            ws1[f"M{r}"] = None
            ws1[f"N{r}"] = None
            ws1[f"O{r}"] = ""
            ws1[f"P{r}"] = f'=IF(B{r}="","",IF(L{r}="Yes","LMS Completed (Ready for Exam)",IF(L{r}="In Progress","LMS In Progress",IF(H{r}="Yes","ADD Attended - Awaiting LMS",IF(H{r}="Rescheduled","ADD Rescheduled",IF(H{r}="Pending","ADD Scheduled",IF(H{r}="No","ADD Missed - Follow Up","Initial Prospect")))))))'
            ws1[f"Q{r}"] = ""
            ws1[f"R{r}"] = None
            ws1[f"S{r}"] = ""

        # Apply cell styling
        for col_letter, _, align, _ in headers:
            cell = ws1[f"{col_letter}{r}"]
            cell.font = font_data
            cell.fill = row_fill
            cell.border = cell_border
            cell.alignment = Alignment(horizontal=align, vertical="center")

            # Number and date formats
            if col_letter in ["F", "G", "I", "K", "M", "R"]:
                cell.number_format = "yyyy-mm-dd"
            elif col_letter == "N":
                cell.number_format = "0.0%"

    # Add Data Validations
    dv_add = DataValidation(type="list", formula1='"Yes,No,Rescheduled,Pending"', allow_blank=True)
    ws1.add_data_validation(dv_add)
    dv_add.add("H9:H65")

    dv_lms = DataValidation(type="list", formula1='"Yes,In Progress,Not Started,Withdrawn"', allow_blank=True)
    ws1.add_data_validation(dv_lms)
    dv_lms.add("L9:L65")

    # Add Conditional Formatting
    fill_green = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    font_green = Font(name="Calibri", size=10, bold=True, color="155724")
    fill_red = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    font_red = Font(name="Calibri", size=10, bold=True, color="721C24")
    fill_amber = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    font_amber = Font(name="Calibri", size=10, bold=True, color="856404")
    fill_blue = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_blue = Font(name="Calibri", size=10, bold=True, color="1E40AF")
    fill_purple = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
    font_purple = Font(name="Calibri", size=10, bold=True, color="5B21B6")
    fill_slate = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_slate = Font(name="Calibri", size=10, color="475569")

    # Col H: Attended ADD
    ws1.conditional_formatting.add("H9:H65", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill_green, font=font_green))
    ws1.conditional_formatting.add("H9:H65", CellIsRule(operator="equal", formula=['"No"'], fill=fill_red, font=font_red))
    ws1.conditional_formatting.add("H9:H65", CellIsRule(operator="equal", formula=['"Rescheduled"'], fill=fill_amber, font=font_amber))
    ws1.conditional_formatting.add("H9:H65", CellIsRule(operator="equal", formula=['"Pending"'], fill=fill_blue, font=font_blue))

    # Col L: Completed LMS
    ws1.conditional_formatting.add("L9:L65", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill_green, font=font_green))
    ws1.conditional_formatting.add("L9:L65", CellIsRule(operator="equal", formula=['"In Progress"'], fill=fill_amber, font=font_amber))
    ws1.conditional_formatting.add("L9:L65", CellIsRule(operator="equal", formula=['"Not Started"'], fill=fill_slate, font=font_slate))
    ws1.conditional_formatting.add("L9:L65", CellIsRule(operator="equal", formula=['"Withdrawn"'], fill=fill_red, font=font_red))

    # Col P: Pipeline Stage
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"LMS Completed (Ready for Exam)"'], fill=fill_green, font=font_green))
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"LMS In Progress"'], fill=fill_amber, font=font_amber))
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"ADD Attended - Awaiting LMS"'], fill=fill_blue, font=font_blue))
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"ADD Scheduled"'], fill=fill_purple, font=font_purple))
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"ADD Missed - Follow Up"'], fill=fill_red, font=font_red))
    ws1.conditional_formatting.add("P9:P65", CellIsRule(operator="equal", formula=['"ADD Rescheduled"'], fill=fill_amber, font=font_amber))

    # Set column widths
    col_widths = {
        "A": 12, "B": 26, "C": 18, "D": 27, "E": 20, "F": 14,
        "G": 18, "H": 15, "I": 19, "J": 32, "K": 17, "L": 17,
        "M": 19, "N": 15, "O": 28, "P": 30, "Q": 32, "R": 18, "S": 35
    }
    for col, width in col_widths.items():
        ws1.column_dimensions[col].width = width


    # -------------------------------------------------------------
    # 2. SHEET 2: Recruitment Dashboard
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Recruitment Dashboard")
    ws2.views.sheetView[0].showGridLines = True

    # Title
    ws2.merge_cells("A1:K1")
    ws2["A1"] = "RECRUITMENT PERFORMANCE DASHBOARD & METRICS"
    ws2["A1"].font = font_title
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:K2")
    ws2["A2"] = "Recruitment Funnel Conversion Analysis, Milestone Breakdown & Recruiter Performance"
    ws2["A2"].font = font_subtitle
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws2.row_dimensions[2].height = 20
    ws2.row_dimensions[3].height = 10

    # Section 1: Funnel Analysis Table
    ws2.merge_cells("A4:D4")
    ws2["A4"] = "RECRUITMENT PIPELINE FUNNEL ANALYSIS"
    ws2["A4"].font = font_sec_hdr
    ws2["A4"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["A4"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    ws2.row_dimensions[4].height = 24

    funnel_headers = [("A", "Funnel Stage", "left"), ("B", "Count", "center"), ("C", "% of Sourced", "center"), ("D", "Conversion / Drop-off Note", "left")]
    for col, title, align in funnel_headers:
        cell = ws2[f"{col}5"]
        cell.value = title
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        cell.border = header_border
    ws2.row_dimensions[5].height = 24

    funnel_rows = [
        (6, "1. Recruits Sourced", "='Recruit Monitoring Tracker'!B5", "=IFERROR(B6/B6,1)", "Total candidates identified in pipeline"),
        (7, "2. Attended ADD", "='Recruit Monitoring Tracker'!F5", "=IFERROR(B7/B6,0)", "Completed Agency Discovery Day seminar"),
        (8, "3. Started LMS", '=COUNTIFS(\'Recruit Monitoring Tracker\'!L9:L65,"<>Not Started",\'Recruit Monitoring Tracker\'!L9:L65,"<>",\'Recruit Monitoring Tracker\'!B9:B65,"<>")', "=IFERROR(B8/B6,0)", "Candidate activated LMS & initiated training"),
        (9, "4. Completed LMS", "='Recruit Monitoring Tracker'!N5", "=IFERROR(B9/B6,0)", "Finished 100% of modules & certified"),
        (10, "5. Ready for Exam", '=COUNTIF(\'Recruit Monitoring Tracker\'!P9:P65, "*Ready for Exam*")', "=IFERROR(B10/B6,0)", "Endorsed for Insurance Commission Licensing"),
    ]

    for r, stage, cnt_formula, pct_formula, note in funnel_rows:
        ws2[f"A{r}"] = stage
        ws2[f"A{r}"].font = font_data_bold
        ws2[f"A{r}"].border = cell_border
        ws2[f"A{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"B{r}"] = cnt_formula
        ws2[f"B{r}"].font = font_data_bold
        ws2[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"B{r}"].border = cell_border
        ws2[f"B{r}"].number_format = "#,##0"
        ws2[f"B{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"C{r}"] = pct_formula
        ws2[f"C{r}"].font = font_data_bold
        ws2[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"C{r}"].border = cell_border
        ws2[f"C{r}"].number_format = "0.0%"
        ws2[f"C{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"D{r}"] = note
        ws2[f"D{r}"].font = font_data
        ws2[f"D{r}"].border = cell_border
        ws2[f"D{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")
        ws2.row_dimensions[r].height = 20

    # Funnel Column Chart
    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Recruitment Funnel Conversion"
    chart1.y_axis.title = "Number of Candidates"
    chart1.x_axis.title = "Pipeline Milestone"
    chart1.height = 9
    chart1.width = 16
    chart1.legend = None

    data_ref = Reference(ws2, min_col=2, min_row=5, max_row=10)
    cats_ref = Reference(ws2, min_col=1, min_row=6, max_row=10)
    chart1.add_data(data_ref, titles_from_data=True)
    chart1.set_categories(cats_ref)
    ws2.add_chart(chart1, "F4")

    # Section 2: Breakdown Tables (Row 12)
    # ADD Breakdown (A12:C18)
    ws2.merge_cells("A12:C12")
    ws2["A12"] = "ADD ATTENDANCE BREAKDOWN"
    ws2["A12"].font = font_sec_hdr
    ws2["A12"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["A12"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    ws2.row_dimensions[12].height = 22

    for col, title, align in [("A", "ADD Status", "left"), ("B", "Count", "center"), ("C", "% Share", "center")]:
        cell = ws2[f"{col}13"]
        cell.value = title
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        cell.border = header_border
    ws2.row_dimensions[13].height = 22

    add_statuses = [
        (14, "Attended (Yes)", '=COUNTIF(\'Recruit Monitoring Tracker\'!H9:H65, "Yes")'),
        (15, "Pending / Scheduled", '=COUNTIF(\'Recruit Monitoring Tracker\'!H9:H65, "Pending")'),
        (16, "Rescheduled", '=COUNTIF(\'Recruit Monitoring Tracker\'!H9:H65, "Rescheduled")'),
        (17, "Missed / No-Show", '=COUNTIF(\'Recruit Monitoring Tracker\'!H9:H65, "No")'),
    ]
    for r, st, formula in add_statuses:
        ws2[f"A{r}"] = st
        ws2[f"A{r}"].font = font_data
        ws2[f"A{r}"].border = cell_border
        ws2[f"A{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"B{r}"] = formula
        ws2[f"B{r}"].font = font_data_bold
        ws2[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"B{r}"].border = cell_border
        ws2[f"B{r}"].number_format = "#,##0"
        ws2[f"B{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"C{r}"] = f"=IFERROR(B{r}/$B$18, 0)"
        ws2[f"C{r}"].font = font_data
        ws2[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"C{r}"].border = cell_border
        ws2[f"C{r}"].number_format = "0.0%"
        ws2[f"C{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")
        ws2.row_dimensions[r].height = 20

    # Total ADD Row
    ws2["A18"] = "Total"
    ws2["A18"].font = font_data_bold
    ws2["A18"].border = total_border
    ws2["B18"] = "=SUM(B14:B17)"
    ws2["B18"].font = font_data_bold
    ws2["B18"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["B18"].border = total_border
    ws2["B18"].number_format = "#,##0"
    ws2["C18"] = "=SUM(C14:C17)"
    ws2["C18"].font = font_data_bold
    ws2["C18"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["C18"].border = total_border
    ws2["C18"].number_format = "0.0%"
    ws2.row_dimensions[18].height = 22

    # LMS Breakdown (D12:F18)
    ws2.merge_cells("D12:F12")
    ws2["D12"] = "LMS COMPLETION BREAKDOWN"
    ws2["D12"].font = font_sec_hdr
    ws2["D12"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["D12"].fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")

    for col, title, align in [("D", "LMS Status", "left"), ("E", "Count", "center"), ("F", "% Share", "center")]:
        cell = ws2[f"{col}13"]
        cell.value = title
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.fill = PatternFill(start_color="14B8A6", end_color="14B8A6", fill_type="solid")
        cell.border = header_border

    lms_statuses = [
        (14, "Completed (Yes)", '=COUNTIF(\'Recruit Monitoring Tracker\'!L9:L65, "Yes")'),
        (15, "In Progress", '=COUNTIF(\'Recruit Monitoring Tracker\'!L9:L65, "In Progress")'),
        (16, "Not Started", '=COUNTIF(\'Recruit Monitoring Tracker\'!L9:L65, "Not Started")'),
        (17, "Withdrawn / Dropped", '=COUNTIF(\'Recruit Monitoring Tracker\'!L9:L65, "Withdrawn")'),
    ]
    for r, st, formula in lms_statuses:
        ws2[f"D{r}"] = st
        ws2[f"D{r}"].font = font_data
        ws2[f"D{r}"].border = cell_border
        ws2[f"D{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"E{r}"] = formula
        ws2[f"E{r}"].font = font_data_bold
        ws2[f"E{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"E{r}"].border = cell_border
        ws2[f"E{r}"].number_format = "#,##0"
        ws2[f"E{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

        ws2[f"F{r}"] = f"=IFERROR(E{r}/$E$18, 0)"
        ws2[f"F{r}"].font = font_data
        ws2[f"F{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"F{r}"].border = cell_border
        ws2[f"F{r}"].number_format = "0.0%"
        ws2[f"F{r}"].fill = PatternFill(start_color="FFFFFF" if r%2==0 else "F8FAFC", end_color="FFFFFF" if r%2==0 else "F8FAFC", fill_type="solid")

    # Total LMS Row
    ws2["D18"] = "Total"
    ws2["D18"].font = font_data_bold
    ws2["D18"].border = total_border
    ws2["E18"] = "=SUM(E14:E17)"
    ws2["E18"].font = font_data_bold
    ws2["E18"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["E18"].border = total_border
    ws2["E18"].number_format = "#,##0"
    ws2["F18"] = "=SUM(F14:F17)"
    ws2["F18"].font = font_data_bold
    ws2["F18"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["F18"].border = total_border
    ws2["F18"].number_format = "0.0%"

    # Pie Chart for LMS Status
    pie_chart = PieChart()
    pie_chart.title = "LMS Completion Breakdown"
    pie_chart.height = 7.5
    pie_chart.width = 11.5
    lms_data_ref = Reference(ws2, min_col=5, min_row=13, max_row=17)
    lms_cats_ref = Reference(ws2, min_col=4, min_row=14, max_row=17)
    pie_chart.add_data(lms_data_ref, titles_from_data=True)
    pie_chart.set_categories(lms_cats_ref)
    ws2.add_chart(pie_chart, "H12")

    # Section 3: Recruiter Performance Table (Row 20)
    ws2.merge_cells("A20:G20")
    ws2["A20"] = "RECRUITER & TEAM PERFORMANCE SUMMARY"
    ws2["A20"].font = font_sec_hdr
    ws2["A20"].alignment = Alignment(horizontal="center", vertical="center")
    ws2["A20"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws2.row_dimensions[20].height = 24

    recruiter_headers = [
        ("A", "Recruiter / Team Leader", "left"),
        ("B", "Total Sourced", "center"),
        ("C", "Attended ADD", "center"),
        ("D", "ADD Rate %", "center"),
        ("E", "Completed LMS", "center"),
        ("F", "LMS Rate %", "center"),
        ("G", "Ready for Exam", "center"),
    ]
    for col, title, align in recruiter_headers:
        cell = ws2[f"{col}21"]
        cell.value = title
        cell.font = font_col_hdr
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.fill = PatternFill(start_color="244B7A", end_color="244B7A", fill_type="solid")
        cell.border = header_border
    ws2.row_dimensions[21].height = 24

    recruiters = ["Sarah Jenkins (UM)", "David Tan", "Elena Gomez", "Michael Chang"]
    for idx, name in enumerate(recruiters, start=22):
        bg = "FFFFFF" if idx%2==0 else "F8FAFC"
        r_fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

        ws2[f"A{idx}"] = name
        ws2[f"A{idx}"].font = font_data_bold
        ws2[f"A{idx}"].border = cell_border
        ws2[f"A{idx}"].fill = r_fill

        ws2[f"B{idx}"] = f'=COUNTIF(\'Recruit Monitoring Tracker\'!E9:E65, "*{name}*")'
        ws2[f"B{idx}"].font = font_data_bold
        ws2[f"B{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"B{idx}"].border = cell_border
        ws2[f"B{idx}"].number_format = "#,##0"
        ws2[f"B{idx}"].fill = r_fill

        ws2[f"C{idx}"] = f'=COUNTIFS(\'Recruit Monitoring Tracker\'!E9:E65, "*{name}*", \'Recruit Monitoring Tracker\'!H9:H65, "Yes")'
        ws2[f"C{idx}"].font = font_data
        ws2[f"C{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"C{idx}"].border = cell_border
        ws2[f"C{idx}"].number_format = "#,##0"
        ws2[f"C{idx}"].fill = r_fill

        ws2[f"D{idx}"] = f'=IFERROR(C{idx}/B{idx}, 0)'
        ws2[f"D{idx}"].font = font_data
        ws2[f"D{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"D{idx}"].border = cell_border
        ws2[f"D{idx}"].number_format = "0.0%"
        ws2[f"D{idx}"].fill = r_fill

        ws2[f"E{idx}"] = f'=COUNTIFS(\'Recruit Monitoring Tracker\'!E9:E65, "*{name}*", \'Recruit Monitoring Tracker\'!L9:L65, "Yes")'
        ws2[f"E{idx}"].font = font_data
        ws2[f"E{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"E{idx}"].border = cell_border
        ws2[f"E{idx}"].number_format = "#,##0"
        ws2[f"E{idx}"].fill = r_fill

        ws2[f"F{idx}"] = f'=IFERROR(E{idx}/C{idx}, 0)'
        ws2[f"F{idx}"].font = font_data
        ws2[f"F{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"F{idx}"].border = cell_border
        ws2[f"F{idx}"].number_format = "0.0%"
        ws2[f"F{idx}"].fill = r_fill

        ws2[f"G{idx}"] = f'=COUNTIFS(\'Recruit Monitoring Tracker\'!E9:E65, "*{name}*", \'Recruit Monitoring Tracker\'!P9:P65, "*Ready for Exam*")'
        ws2[f"G{idx}"].font = font_data_bold
        ws2[f"G{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"G{idx}"].border = cell_border
        ws2[f"G{idx}"].number_format = "#,##0"
        ws2[f"G{idx}"].fill = r_fill
        ws2.row_dimensions[idx].height = 20

    # Total Recruiter Row
    tot_r = 26
    ws2[f"A{tot_r}"] = "Total Agency"
    ws2[f"A{tot_r}"].font = font_data_bold
    ws2[f"A{tot_r}"].border = total_border

    for col in ["B", "C", "E", "G"]:
        ws2[f"{col}{tot_r}"] = f"=SUM({col}22:{col}25)"
        ws2[f"{col}{tot_r}"].font = font_data_bold
        ws2[f"{col}{tot_r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2[f"{col}{tot_r}"].border = total_border
        ws2[f"{col}{tot_r}"].number_format = "#,##0"

    ws2[f"D{tot_r}"] = f"=IFERROR(C{tot_r}/B{tot_r}, 0)"
    ws2[f"D{tot_r}"].font = font_data_bold
    ws2[f"D{tot_r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws2[f"D{tot_r}"].border = total_border
    ws2[f"D{tot_r}"].number_format = "0.0%"

    ws2[f"F{tot_r}"] = f"=IFERROR(E{tot_r}/C{tot_r}, 0)"
    ws2[f"F{tot_r}"].font = font_data_bold
    ws2[f"F{tot_r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws2[f"F{tot_r}"].border = total_border
    ws2[f"F{tot_r}"].number_format = "0.0%"
    ws2.row_dimensions[tot_r].height = 22

    # Column widths for Dashboard
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["D"].width = 24
    ws2.column_dimensions["E"].width = 16
    ws2.column_dimensions["F"].width = 18
    ws2.column_dimensions["G"].width = 18
    ws2.column_dimensions["H"].width = 16
    ws2.column_dimensions["I"].width = 16
    ws2.column_dimensions["J"].width = 16


    # -------------------------------------------------------------
    # 3. SHEET 3: Dropdown Lists & Settings
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Dropdowns & Settings")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:H1")
    ws3["A1"] = "SYSTEM DROPDOWN LISTS & CONFIGURATION REFERENCE"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws3.row_dimensions[1].height = 36

    ws3.merge_cells("A2:H2")
    ws3["A2"] = "Reference options for tracker dropdown menus, funnel stages, and agency recruiters"
    ws3["A2"].font = font_subtitle
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws3.row_dimensions[2].height = 20

    # Table 1: ADD Status Options (A4:B9)
    ws3.merge_cells("A4:B4")
    ws3["A4"] = "ADD ATTENDANCE OPTIONS"
    ws3["A4"].font = font_sec_hdr
    ws3["A4"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["A4"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    ws3.row_dimensions[4].height = 22

    ws3["A5"] = "Option Value"
    ws3["A5"].font = font_col_hdr
    ws3["A5"].fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    ws3["A5"].border = header_border

    ws3["B5"] = "Definition & Action Trigger"
    ws3["B5"].font = font_col_hdr
    ws3["B5"].fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    ws3["B5"].border = header_border
    ws3.row_dimensions[5].height = 22

    add_defs = [
        ("Yes", "Candidate attended ADD orientation session. Action: Issue LMS credentials."),
        ("No", "Candidate missed session (no-show). Action: Call back & re-invite to next batch."),
        ("Rescheduled", "Candidate requested new schedule prior to session. Action: Confirm new batch date."),
        ("Pending", "Candidate confirmed attendance for upcoming ADD batch. Action: Send reminder."),
    ]
    for idx, (opt, desc) in enumerate(add_defs, start=6):
        ws3[f"A{idx}"] = opt
        ws3[f"A{idx}"].font = font_data_bold
        ws3[f"A{idx}"].border = cell_border
        ws3[f"B{idx}"] = desc
        ws3[f"B{idx}"].font = font_data
        ws3[f"B{idx}"].border = cell_border
        ws3.row_dimensions[idx].height = 20

    # Table 2: LMS Status Options (D4:E9)
    ws3.merge_cells("D4:E4")
    ws3["D4"] = "LMS COMPLETION OPTIONS"
    ws3["D4"].font = font_sec_hdr
    ws3["D4"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["D4"].fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")

    ws3["D5"] = "Option Value"
    ws3["D5"].font = font_col_hdr
    ws3["D5"].fill = PatternFill(start_color="14B8A6", end_color="14B8A6", fill_type="solid")
    ws3["D5"].border = header_border

    ws3["E5"] = "Definition & Action Trigger"
    ws3["E5"].font = font_col_hdr
    ws3["E5"].fill = PatternFill(start_color="14B8A6", end_color="14B8A6", fill_type="solid")
    ws3["E5"].border = header_border

    lms_defs = [
        ("Yes", "100% completed & certified. Action: Endorse for Licensing Exam."),
        ("In Progress", "Currently taking modules (20%-90%). Action: Check progress & nudge completion."),
        ("Not Started", "Account created or invited, 0% completed. Action: Resend credentials & setup call."),
        ("Withdrawn", "Candidate discontinued application. Action: Archive record or note reason."),
    ]
    for idx, (opt, desc) in enumerate(lms_defs, start=6):
        ws3[f"D{idx}"] = opt
        ws3[f"D{idx}"].font = font_data_bold
        ws3[f"D{idx}"].border = cell_border
        ws3[f"E{idx}"] = desc
        ws3[f"E{idx}"].font = font_data
        ws3[f"E{idx}"].border = cell_border

    # Table 3: Recruiter List (G4:H9)
    ws3.merge_cells("G4:H4")
    ws3["G4"] = "AGENCY RECRUITERS / LEADERS"
    ws3["G4"].font = font_sec_hdr
    ws3["G4"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["G4"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")

    ws3["G5"] = "Recruiter Name"
    ws3["G5"].font = font_col_hdr
    ws3["G5"].fill = PatternFill(start_color="244B7A", end_color="244B7A", fill_type="solid")
    ws3["G5"].border = header_border

    ws3["H5"] = "Role / Designation"
    ws3["H5"].font = font_col_hdr
    ws3["H5"].fill = PatternFill(start_color="244B7A", end_color="244B7A", fill_type="solid")
    ws3["H5"].border = header_border

    recruiters_cfg = [
        ("Sarah Jenkins (UM)", "Unit Manager / Agency Leader"),
        ("David Tan", "Senior Financial Advisor / Recruiter"),
        ("Elena Gomez", "Agency Recruitment Specialist"),
        ("Michael Chang", "Financial Advisor / Associate"),
    ]
    for idx, (name, role) in enumerate(recruiters_cfg, start=6):
        ws3[f"G{idx}"] = name
        ws3[f"G{idx}"].font = font_data_bold
        ws3[f"G{idx}"].border = cell_border
        ws3[f"H{idx}"] = role
        ws3[f"H{idx}"].font = font_data
        ws3[f"H{idx}"].border = cell_border

    # Table 4: Pipeline Stages Explained (A12:E19)
    ws3.merge_cells("A12:E12")
    ws3["A12"] = "AUTOMATED PIPELINE STAGES & LOGIC EXPLAINED"
    ws3["A12"].font = font_sec_hdr
    ws3["A12"].alignment = Alignment(horizontal="center", vertical="center")
    ws3["A12"].fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws3.row_dimensions[12].height = 22

    ws3["A13"] = "Pipeline Stage"
    ws3["A13"].font = font_col_hdr
    ws3["A13"].fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    ws3["A13"].border = header_border

    ws3.merge_cells("B13:E13")
    ws3["B13"] = "Trigger Condition in Workbook Formulas"
    ws3["B13"].font = font_col_hdr
    ws3["B13"].fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    ws3["B13"].border = header_border
    ws3.row_dimensions[13].height = 22

    pipeline_stages = [
        ("LMS Completed (Ready for Exam)", "Completed LMS? = 'Yes' (Highest milestone achieved; endorse to regulatory exam)"),
        ("LMS In Progress", "Completed LMS? = 'In Progress' (Recruit is actively doing training modules)"),
        ("ADD Attended - Awaiting LMS", "Attended ADD? = 'Yes' and LMS not yet started (Ready to receive credentials)"),
        ("ADD Scheduled", "Attended ADD? = 'Pending' (Invite sent, scheduled for orientation)"),
        ("ADD Rescheduled", "Attended ADD? = 'Rescheduled' (Candidate requested future batch)"),
        ("ADD Missed - Follow Up", "Attended ADD? = 'No' (Candidate did not attend; requires urgent re-engagement)"),
    ]
    for idx, (stage, cond) in enumerate(pipeline_stages, start=14):
        ws3[f"A{idx}"] = stage
        ws3[f"A{idx}"].font = font_data_bold
        ws3[f"A{idx}"].border = cell_border
        ws3.merge_cells(f"B{idx}:E{idx}")
        ws3[f"B{idx}"] = cond
        ws3[f"B{idx}"].font = font_data
        for c in ["B", "C", "D", "E"]:
            ws3[f"{c}{idx}"].border = cell_border
        ws3.row_dimensions[idx].height = 20

    # Column widths for Settings
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 50
    ws3.column_dimensions["C"].width = 6
    ws3.column_dimensions["D"].width = 22
    ws3.column_dimensions["E"].width = 50
    ws3.column_dimensions["F"].width = 6
    ws3.column_dimensions["G"].width = 26
    ws3.column_dimensions["H"].width = 35


    # -------------------------------------------------------------
    # 4. SHEET 4: Quick User Guide
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="User Guide & Workflow")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    ws4["A1"] = "RECRUIT MONITORING TRACKER - USER GUIDE & BEST PRACTICES"
    ws4["A1"].font = font_title
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws4["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws4.row_dimensions[1].height = 36

    ws4.merge_cells("A2:G2")
    ws4["A2"] = "Step-by-step operating guide for tracking candidate journey from ADD attendance to LMS completion"
    ws4["A2"].font = font_subtitle
    ws4["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws4["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws4.row_dimensions[2].height = 20

    guide_sections = [
        ("1. OVERVIEW & OBJECTIVE", [
            ("Purpose", "This workbook is customized for agency leaders and recruiters to track candidates systematically through the two critical recruitment hurdles: (1) Attending ADD (Agency / Agent Discovery Day) and (2) Completing LMS (Learning Management System pre-licensing training)."),
            ("Design", "Automated KPI summary cards, formula-driven pipeline stage classification, built-in data validation dropdowns, visual color highlighting, and a full executive analytics dashboard."),
        ]),
        ("2. STEP-BY-STEP RECRUITMENT WORKFLOW", [
            ("Step 1: Input Recruit Profile", "Type the recruit's Complete Name in Column B (Last Name, First Name M.I.). The Recruit ID (Column A) will automatically generate! Enter contact number, email, and the sourcing recruiter."),
            ("Step 2: Monitor ADD Attendance", "Enter the ADD Scheduled Date in Col G. When the session concludes, set 'Attended ADD?' (Col H) to 'Yes', 'No', or 'Rescheduled' using the dropdown menu. Record the actual date attended in Col I and brief feedback in Col J."),
            ("Step 3: Track LMS Completion", "Once the recruit passes ADD, issue LMS access (Col K). As the candidate studies, update 'Completed LMS?' (Col L) to 'In Progress' or 'Yes'. Enter their module progress % in Col N (e.g. 50%, 100%) and certification details in Col O."),
            ("Step 4: Execute Next Actions", "The workbook will automatically update 'Current Pipeline Stage' (Col P). Refer to this column to see who is 'Ready for Exam', who is 'LMS In Progress', and who missed ADD. Enter the Next Action (Col Q) and Follow-up Date (Col R)."),
        ]),
        ("3. KEY METRICS & FORMULA AUTOMATION", [
            ("Attendance Rate %", "Calculates Attended ADD / Total Recruits Sourced. Benchmark target is >= 70%."),
            ("Completion Rate %", "Calculates Completed LMS / Attended ADD. Benchmark target is >= 60%."),
            ("Funnel Conversion %", "Calculates Completed LMS / Total Recruits Sourced. Benchmark target is >= 40%."),
            ("Adding New Recruits", "Rows 15 to 65 are pre-formatted with all dropdowns, formulas, and borders. Simply start typing in Column B (Complete Name), and the entire row will activate automatically!"),
        ]),
        ("4. BEST PRACTICES FOR RECRUITING SUCCESS", [
            ("The 24-Hour Rule", "Reach out to ADD attendees within 24 hours while interest is fresh to set up LMS credentials."),
            ("LMS Momentum Window", "Encourage recruits to complete the 6 LMS modules within 7 to 10 days before enthusiasm wanes."),
            ("Re-scheduling Protocol", "For no-shows ('No'), call within 48 hours to offer 1-on-1 career briefing or reschedule to the next ADD batch."),
            ("Mock Exam Review", "Once LMS is Completed ('Yes'), immediately schedule a practice mock test before filing for the licensing exam."),
        ])
    ]

    curr_row = 4
    for sec_title, items in guide_sections:
        ws4.merge_cells(f"A{curr_row}:G{curr_row}")
        ws4[f"A{curr_row}"] = sec_title
        ws4[f"A{curr_row}"].font = font_sec_hdr
        ws4[f"A{curr_row}"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        ws4[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws4.row_dimensions[curr_row].height = 24
        curr_row += 1

        for label, text in items:
            ws4[f"A{curr_row}"] = label
            ws4[f"A{curr_row}"].font = font_data_bold
            ws4[f"A{curr_row}"].fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            ws4[f"A{curr_row}"].border = cell_border
            ws4[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="top")

            ws4.merge_cells(f"B{curr_row}:G{curr_row}")
            ws4[f"B{curr_row}"] = text
            ws4[f"B{curr_row}"].font = font_data
            ws4[f"B{curr_row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            for c in ["B", "C", "D", "E", "F", "G"]:
                ws4[f"{c}{curr_row}"].border = cell_border
                ws4[f"{c}{curr_row}"].fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            ws4.row_dimensions[curr_row].height = 36
            curr_row += 1

        curr_row += 1 # blank row between sections

    ws4.column_dimensions["A"].width = 24
    ws4.column_dimensions["B"].width = 20
    ws4.column_dimensions["C"].width = 20
    ws4.column_dimensions["D"].width = 20
    ws4.column_dimensions["E"].width = 20
    ws4.column_dimensions["F"].width = 20
    ws4.column_dimensions["G"].width = 20

    # Save workbook
    wb.save(filename)
    print(f"Workbook successfully saved to: {filename}")

if __name__ == "__main__":
    create_recruit_workbook()
