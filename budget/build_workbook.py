#!/usr/bin/env python3
"""
Build the editable edition of the MONTHLY BUDGET PLANNER.

An .xlsx workbook that opens in Excel, Numbers, LibreOffice and Google Sheets.
Every total, difference and percentage is a live formula, so the workbook adds
up as you type. Styling mirrors the printed planner: cream fields, sage
headers, serif-style titles, no hard-coded numbers anywhere.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import (Alignment, Border, Font, PatternFill, Side)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import content as C

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")

CREAM = "FDFBF6"
PANEL = "F5F0E6"
PANEL2 = "EFEADF"
SAGE = "8B9E86"
SAGE_D = "6E8169"
BROWN = "4E4238"
GREY = "8A7F73"
ROSE = "C29A94"
LINE = "D8CDBA"
WHITE = "FFFFFF"

TITLE_F = Font(name="Georgia", size=20, bold=True, color=BROWN)
SUB_F = Font(name="Georgia", size=11, italic=True, color=GREY)
HEAD_F = Font(name="Calibri", size=10, bold=True, color=WHITE)
LBL_F = Font(name="Calibri", size=10, bold=True, color=BROWN)
BODY_F = Font(name="Calibri", size=10, color=BROWN)
HINT_F = Font(name="Calibri", size=9, italic=True, color=GREY)
TOT_F = Font(name="Calibri", size=11, bold=True, color=SAGE_D)
SEC_F = Font(name="Calibri", size=10, bold=True, color=SAGE_D)

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_H = PatternFill("solid", fgColor=SAGE)
FILL_P = PatternFill("solid", fgColor=PANEL)
FILL_P2 = PatternFill("solid", fgColor=PANEL2)
FILL_IN = PatternFill("solid", fgColor=WHITE)
FILL_BG = PatternFill("solid", fgColor=CREAM)

MONEY = '"$"#,##0.00'
PCT = "0.0%"

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def sheet(wb, name, widths, tab=SAGE):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(1, 90):
        for c in range(1, len(widths) + 2):
            ws.cell(row=r, column=c).fill = FILL_BG
    return ws


def title(ws, text, ncols, sub=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(row=1, column=1, value=text)
    t.font = TITLE_F
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        s = ws.cell(row=2, column=1, value=sub)
        s.font = SUB_F
        s.alignment = CENTER
        ws.row_dimensions[2].height = 18
        return 4
    return 3


def header_row(ws, r, cols, start=1):
    for i, name in enumerate(cols):
        c = ws.cell(row=r, column=start + i, value=name)
        c.font = HEAD_F
        c.fill = FILL_H
        c.alignment = CENTER
        c.border = BOX
    ws.row_dimensions[r].height = 22


def input_cells(ws, r0, r1, c0, c1, fmt=None, zebra=True):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = BODY_F
            cell.fill = FILL_P if (zebra and (r - r0) % 2 == 1) else FILL_IN
            if fmt:
                cell.number_format = fmt
            cell.alignment = CENTER
        ws.row_dimensions[r].height = 19


def label_value(ws, r, label, col=1, vcol=2, span=3, fmt=None):
    c = ws.cell(row=r, column=col, value=label)
    c.font = LBL_F
    c.alignment = LEFT
    ws.merge_cells(start_row=r, start_column=vcol, end_row=r,
                   end_column=vcol + span - 1)
    v = ws.cell(row=r, column=vcol)
    v.fill = FILL_IN
    v.border = BOX
    v.font = BODY_F
    if fmt:
        v.number_format = fmt
    ws.row_dimensions[r].height = 22


def band(ws, r, text, ncols, fmt=None, formula=None, fill=FILL_P2):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols - 1)
    c = ws.cell(row=r, column=1, value=text)
    c.font = TOT_F
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.fill = fill
    for i in range(1, ncols):
        ws.cell(row=r, column=i).fill = fill
    v = ws.cell(row=r, column=ncols, value=formula)
    v.font = TOT_F
    v.fill = fill
    v.border = BOX
    v.alignment = CENTER
    if fmt:
        v.number_format = fmt
    ws.row_dimensions[r].height = 24
    return v


def note(ws, r, text, ncols):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(row=r, column=1, value=text)
    c.font = HINT_F
    c.alignment = CENTER


def yesno(ws, ref, options='"Yes,No"'):
    dv = DataValidation(type="list", formula1=options, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(ref)


# ------------------------------------------------------------------ sheets
def s_start(wb):
    ws = sheet(wb, "Start Here", [3, 26, 30, 30, 20], tab=ROSE)
    ws.merge_cells("B2:E2")
    t = ws["B2"]
    t.value = C.TITLE
    t.font = Font(name="Georgia", size=24, bold=True, color=BROWN)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 40
    ws.merge_cells("B3:E3")
    ws["B3"].value = C.SUBTITLE
    ws["B3"].font = Font(name="Georgia", size=12, italic=True, color=GREY)
    ws["B3"].alignment = CENTER
    ws.merge_cells("B4:E4")
    ws["B4"].value = C.TAGLINE
    ws["B4"].font = Font(name="Calibri", size=11, bold=True, color=SAGE_D)
    ws["B4"].alignment = CENTER

    r = 6
    for f in ["Name:", "Year:", "My financial focus this year:"]:
        label_value(ws, r, f, col=2, vcol=3, span=3)
        r += 2

    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=2, value=C.BELONGS["quote"]).font = SUB_F
    ws.cell(row=r, column=2).alignment = CENTER
    r += 3

    ws.cell(row=r, column=2, value="HOW TO USE THIS WORKBOOK").font = SEC_F
    r += 1
    steps = C.HOWTO["steps"] + [
        "Every green figure is calculated for you \u2014 leave those alone.",
        "White boxes are yours to type in.",
        "Duplicate the monthly tabs to start a new month.",
    ]
    for i, s in enumerate(steps, 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        fl = FILL_P if i % 2 else FILL_IN
        c = ws.cell(row=r, column=2, value=f"{i}.   {s}")
        c.font = BODY_F
        c.alignment = LEFT
        for cc in range(2, 6):
            ws.cell(row=r, column=cc).fill = fl
        ws.row_dimensions[r].height = 20
        r += 1

    r += 2
    ws.merge_cells(start_row=r, start_column=2, end_row=r + 3, end_column=5)
    c = ws.cell(row=r, column=2,
                value="This workbook is a personal organizing tool. It records "
                      "the numbers you enter and adds them up for you. It does "
                      "not provide financial, investment, tax or legal advice.")
    c.font = HINT_F
    c.alignment = WRAP
    c.fill = FILL_P2
    return ws


def s_year(wb):
    ws = sheet(wb, "Year at a Glance", [4, 16, 16, 16, 16, 16, 34])
    r = title(ws, C.YEAR_VIEW["title"], 7)
    cols = ["MONTH", "INCOME", "EXPENSES", "SAVINGS", "LEFT OVER", "NOTES"]
    header_row(ws, r, cols, start=2)
    r0 = r + 1
    for i, m in enumerate(C.MONTHS):
        rr = r0 + i
        c = ws.cell(row=rr, column=2, value=m)
        c.font = LBL_F
        c.alignment = LEFT
        c.border = BOX
        c.fill = FILL_P if i % 2 else FILL_IN
        input_cells(ws, rr, rr, 3, 5, MONEY, zebra=False)
        for cc in range(3, 6):
            ws.cell(row=rr, column=cc).fill = FILL_P if i % 2 else FILL_IN
        f = ws.cell(row=rr, column=6,
                    value=f"=IF(COUNT(C{rr}:E{rr})=0,\"\",C{rr}-D{rr}-E{rr})")
        f.number_format = MONEY
        f.font = TOT_F
        f.border = BOX
        f.alignment = CENTER
        f.fill = FILL_P2
        n = ws.cell(row=rr, column=7)
        n.border = BOX
        n.fill = FILL_P if i % 2 else FILL_IN
        n.alignment = LEFT
        ws.row_dimensions[rr].height = 21
    r1 = r0 + 11
    tr = r1 + 1
    ws.cell(row=tr, column=2, value="TOTAL").font = TOT_F
    ws.cell(row=tr, column=2).fill = FILL_P2
    ws.cell(row=tr, column=2).border = BOX
    for cc in range(3, 7):
        L = get_column_letter(cc)
        c = ws.cell(row=tr, column=cc, value=f"=SUM({L}{r0}:{L}{r1})")
        c.number_format = MONEY
        c.font = TOT_F
        c.fill = FILL_P2
        c.border = BOX
        c.alignment = CENTER
    ws.cell(row=tr, column=7).fill = FILL_P2
    ws.cell(row=tr, column=7).border = BOX
    note(ws, tr + 2, "Totals update automatically as you fill in each month.", 7)
    return ws


def s_income(wb):
    ws = sheet(wb, "Income", [4, 14, 30, 16, 16, 28])
    r = title(ws, C.INCOME["title"], 6)
    header_row(ws, r, C.INCOME["cols"], start=2)
    r0 = r + 1
    n = 16
    input_cells(ws, r0, r0 + n - 1, 2, 6)
    for rr in range(r0, r0 + n):
        ws.cell(row=rr, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=rr, column=3).alignment = LEFT
        ws.cell(row=rr, column=4).number_format = MONEY
        ws.cell(row=rr, column=5).number_format = MONEY
        ws.cell(row=rr, column=6).alignment = LEFT
    r1 = r0 + n - 1
    band(ws, r1 + 1, C.INCOME["totals"][0], 6,
         MONEY, f"=SUM(D{r0}:D{r1})")
    band(ws, r1 + 2, C.INCOME["totals"][1], 6,
         MONEY, f"=SUM(E{r0}:E{r1})")
    band(ws, r1 + 3, "DIFFERENCE", 6, MONEY,
         f"=SUM(E{r0}:E{r1})-SUM(D{r0}:D{r1})")
    ws.defined_names  # noqa
    return ws, f"'Income'!E{r1 + 2}"


def s_fixed(wb):
    ws = sheet(wb, "Fixed Expenses", [4, 28, 14, 16, 16, 12])
    r = title(ws, C.FIXED["title"], 6,
              "Example labels are shown in grey \u2014 replace or add your own.")
    header_row(ws, r, C.FIXED["cols"], start=2)
    r0 = r + 1
    n = 16
    input_cells(ws, r0, r0 + n - 1, 2, 6)
    for i, rr in enumerate(range(r0, r0 + n)):
        ws.cell(row=rr, column=2).alignment = LEFT
        if i < len(C.FIXED["examples"]):
            c = ws.cell(row=rr, column=2, value=C.FIXED["examples"][i])
            c.font = Font(name="Calibri", size=10, color="B9AFA1")
            c.alignment = LEFT
        ws.cell(row=rr, column=4).number_format = MONEY
        ws.cell(row=rr, column=5).number_format = MONEY
    r1 = r0 + n - 1
    yesno(ws, f"F{r0}:F{r1}")
    band(ws, r1 + 1, "TOTAL BUDGETED", 6, MONEY, f"=SUM(D{r0}:D{r1})")
    band(ws, r1 + 2, "TOTAL ACTUAL", 6, MONEY, f"=SUM(E{r0}:E{r1})")
    return ws, f"'Fixed Expenses'!E{r1 + 2}"


def s_variable(wb):
    ws = sheet(wb, "Variable Expenses", [4, 26, 18, 18, 18])
    r = title(ws, C.VARIABLE["title"], 5)
    header_row(ws, r, ["CATEGORY"] + C.VARIABLE["fields"], start=2)
    r0 = r + 1
    cats = C.VARIABLE["categories"]
    for i, cat in enumerate(cats):
        rr = r0 + i
        c = ws.cell(row=rr, column=2, value=cat)
        c.font = LBL_F
        c.alignment = LEFT
        c.border = BOX
        c.fill = FILL_P if i % 2 else FILL_IN
        for cc in (3, 4):
            cell = ws.cell(row=rr, column=cc)
            cell.border = BOX
            cell.number_format = MONEY
            cell.alignment = CENTER
            cell.fill = FILL_P if i % 2 else FILL_IN
        d = ws.cell(row=rr, column=5,
                    value=f"=IF(COUNT(C{rr}:D{rr})=0,\"\",C{rr}-D{rr})")
        d.number_format = MONEY
        d.font = TOT_F
        d.border = BOX
        d.alignment = CENTER
        d.fill = FILL_P2
        ws.row_dimensions[rr].height = 22
    r1 = r0 + len(cats) - 1
    band(ws, r1 + 1, "TOTAL", 5, MONEY, f"=SUM(D{r0}:D{r1})")
    note(ws, r1 + 3, "Difference shows budgeted minus actual. A positive "
                     "number means you stayed under.", 5)
    return ws, f"'Variable Expenses'!D{r1 + 1}"


def s_budget(wb, income_ref):
    ws = sheet(wb, "Monthly Budget", [4, 26, 18, 18, 18])
    r = title(ws, C.BUDGET["title"], 5)
    header_row(ws, r, C.BUDGET["cols"], start=2)
    r0 = r + 1
    secs = C.BUDGET["sections"]
    for i, s in enumerate(secs):
        rr = r0 + i
        c = ws.cell(row=rr, column=2, value=s)
        c.font = LBL_F
        c.alignment = LEFT
        c.border = BOX
        c.fill = FILL_P if i % 2 else FILL_IN
        for cc in (3, 4):
            cell = ws.cell(row=rr, column=cc)
            cell.border = BOX
            cell.number_format = MONEY
            cell.alignment = CENTER
            cell.fill = FILL_P if i % 2 else FILL_IN
        d = ws.cell(row=rr, column=5,
                    value=f"=IF(COUNT(C{rr}:D{rr})=0,\"\",C{rr}-D{rr})")
        d.number_format = MONEY
        d.font = TOT_F
        d.border = BOX
        d.alignment = CENTER
        d.fill = FILL_P2
        ws.row_dimensions[rr].height = 22
    r1 = r0 + len(secs) - 1
    tr = r1 + 1
    ws.cell(row=tr, column=2, value=C.BUDGET["total"]).font = TOT_F
    ws.cell(row=tr, column=2).fill = FILL_P2
    ws.cell(row=tr, column=2).border = BOX
    for cc in (3, 4, 5):
        L = get_column_letter(cc)
        c = ws.cell(row=tr, column=cc, value=f"=SUM({L}{r0}:{L}{r1})")
        c.number_format = MONEY
        c.font = TOT_F
        c.fill = FILL_P2
        c.border = BOX
        c.alignment = CENTER
    ws.row_dimensions[tr].height = 24
    note(ws, tr + 2, "Difference shows planned minus actual for each category.", 5)
    return ws, r0, r1


def s_bills(wb):
    ws = sheet(wb, "Bills", [4, 26, 14, 16, 16, 14])
    r = title(ws, C.BILLS["title"], 6, C.BILLS["note"])
    header_row(ws, r, C.BILLS["cols"], start=2)
    r0 = r + 1
    n = 20
    input_cells(ws, r0, r0 + n - 1, 2, 6)
    for rr in range(r0, r0 + n):
        ws.cell(row=rr, column=2).alignment = LEFT
        ws.cell(row=rr, column=4).number_format = MONEY
    r1 = r0 + n - 1
    yesno(ws, f"F{r0}:F{r1}", '"Paid,Due,Overdue"')
    band(ws, r1 + 1, "TOTAL BILLS", 6, MONEY, f"=SUM(D{r0}:D{r1})")
    band(ws, r1 + 2, "NUMBER MARKED PAID", 6, "0",
         f'=COUNTIF(F{r0}:F{r1},"Paid")')
    return ws


def s_subs(wb):
    ws = sheet(wb, "Subscriptions", [4, 24, 16, 16, 16, 26])
    r = title(ws, C.SUBS["title"], 6, C.SUBS["prompts"][0])
    header_row(ws, r, C.SUBS["cols"], start=2)
    r0 = r + 1
    n = 16
    input_cells(ws, r0, r0 + n - 1, 2, 6)
    for rr in range(r0, r0 + n):
        ws.cell(row=rr, column=2).alignment = LEFT
        ws.cell(row=rr, column=3).number_format = MONEY
        ws.cell(row=rr, column=6).alignment = LEFT
    r1 = r0 + n - 1
    yesno(ws, f"E{r0}:E{r1}", '"Keep,Cancel"')
    band(ws, r1 + 1, "TOTAL PER MONTH", 6, MONEY, f"=SUM(C{r0}:C{r1})")
    band(ws, r1 + 2, "TOTAL PER YEAR", 6, MONEY, f"=SUM(C{r0}:C{r1})*12")
    band(ws, r1 + 3, "MONTHLY COST OF THE ONES MARKED KEEP", 6, MONEY,
         f'=SUMIF(E{r0}:E{r1},"Keep",C{r0}:C{r1})')
    band(ws, r1 + 4, "MONTHLY SAVING IF YOU CANCEL THE REST", 6, MONEY,
         f'=SUMIF(E{r0}:E{r1},"Cancel",C{r0}:C{r1})')
    note(ws, r1 + 6, C.SUBS["prompts"][1], 6)
    return ws


def s_savings(wb):
    ws = sheet(wb, "Savings Goals", [4, 26, 18, 18, 18, 18])
    r = title(ws, C.SAVINGS_GOALS["title"], 6)
    header_row(ws, r, ["GOAL", "TARGET AMOUNT", "CURRENT SAVINGS",
                       "DEADLINE", "PROGRESS"], start=2)
    r0 = r + 1
    n = 6
    input_cells(ws, r0, r0 + n - 1, 2, 5)
    for i, rr in enumerate(range(r0, r0 + n)):
        ws.cell(row=rr, column=2).alignment = LEFT
        ws.cell(row=rr, column=3).number_format = MONEY
        ws.cell(row=rr, column=4).number_format = MONEY
        p = ws.cell(row=rr, column=6,
                    value=f'=IF(N(C{rr})=0,"",D{rr}/C{rr})')
        p.number_format = PCT
        p.font = TOT_F
        p.border = BOX
        p.alignment = CENTER
        p.fill = FILL_P2
        ws.row_dimensions[rr].height = 21
    r1 = r0 + n - 1
    band(ws, r1 + 1, "TOTAL SAVED TOWARD GOALS", 6, MONEY,
         f"=SUM(D{r0}:D{r1})")
    band(ws, r1 + 2, "TOTAL STILL TO GO", 6, MONEY,
         f"=MAX(0,SUM(C{r0}:C{r1})-SUM(D{r0}:D{r1}))")

    er = r1 + 4
    ws.merge_cells(start_row=er, start_column=2, end_row=er, end_column=6)
    c = ws.cell(row=er, column=2, value=C.EMERGENCY["title"])
    c.font = Font(name="Georgia", size=14, bold=True, color=BROWN)
    c.alignment = CENTER
    er += 1
    for p in C.EMERGENCY["prompts"]:
        label_value(ws, er, p, col=2, vcol=4, span=2, fmt=MONEY)
        er += 1
    pr = ws.cell(row=er, column=2, value="Progress toward target:")
    pr.font = LBL_F
    ws.merge_cells(start_row=er, start_column=4, end_row=er, end_column=5)
    f = ws.cell(row=er, column=4,
                value=f'=IF(N(D{er - 2})=0,"",D{er - 3}/D{er - 2})')
    f.number_format = PCT
    f.font = TOT_F
    f.fill = FILL_P2
    f.border = BOX
    f.alignment = CENTER
    note(ws, er + 2, C.EMERGENCY["quote"], 6)
    return ws


def s_challenge(wb):
    ws = sheet(wb, "Savings Challenge", [4, 10, 20, 20, 6, 10, 20, 20])
    r = title(ws, C.CHALLENGE["title"], 8)
    label_value(ws, r, C.CHALLENGE["target"], col=2, vcol=3, span=2, fmt=MONEY)
    target_row = r
    r += 2
    header_row(ws, r, C.CHALLENGE["cols"], start=2)
    header_row(ws, r, C.CHALLENGE["cols"], start=6)
    r0 = r + 1
    for col_i, base in enumerate((2, 6)):
        for i in range(15):
            rr = r0 + i
            day = col_i * 15 + i + 1
            d = ws.cell(row=rr, column=base, value=day)
            d.font = LBL_F
            d.border = BOX
            d.alignment = CENTER
            d.fill = FILL_P if i % 2 else FILL_IN
            a = ws.cell(row=rr, column=base + 1)
            a.border = BOX
            a.number_format = MONEY
            a.alignment = CENTER
            a.fill = FILL_P if i % 2 else FILL_IN
            AL = get_column_letter(base + 1)
            if col_i == 0:
                run = (f"=IF(COUNT(${AL}${r0}:{AL}{rr})=0,\"\","
                       f"SUM(${AL}${r0}:{AL}{rr}))")
            else:
                pl = get_column_letter(3)
                run = (f"=IF(COUNT($C${r0}:$C${r0 + 14},"
                       f"${AL}${r0}:{AL}{rr})=0,\"\","
                       f"SUM($C${r0}:$C${r0 + 14})+SUM(${AL}${r0}:{AL}{rr}))")
            t = ws.cell(row=rr, column=base + 2, value=run)
            t.number_format = MONEY
            t.font = TOT_F
            t.border = BOX
            t.alignment = CENTER
            t.fill = FILL_P2
            ws.row_dimensions[rr].height = 19
    r1 = r0 + 14
    tot = f"=SUM(C{r0}:C{r1})+SUM(G{r0}:G{r1})"
    band(ws, r1 + 2, "TOTAL SAVED OVER 30 DAYS", 8, MONEY, tot)
    band(ws, r1 + 3, "AMOUNT LEFT TO REACH YOUR TARGET", 8, MONEY,
         f"=MAX(0,N(C{target_row})-(SUM(C{r0}:C{r1})+SUM(G{r0}:G{r1})))")
    return ws


def s_debt(wb):
    ws = sheet(wb, "Debt", [4, 22, 16, 16, 16, 14, 14, 16])
    r = title(ws, C.DEBT["title"], 8)
    header_row(ws, r, ["DEBT", "STARTING BALANCE", "CURRENT BALANCE",
                       "MINIMUM PAYMENT", "EXTRA PAYMENT", "DUE DATE",
                       "PAID OFF SO FAR"], start=2)
    r0 = r + 1
    n = 12
    input_cells(ws, r0, r0 + n - 1, 2, 7)
    for rr in range(r0, r0 + n):
        ws.cell(row=rr, column=2).alignment = LEFT
        for cc in (3, 4, 5, 6):
            ws.cell(row=rr, column=cc).number_format = MONEY
        p = ws.cell(row=rr, column=8,
                    value=f'=IF(COUNT(C{rr}:D{rr})<2,"",C{rr}-D{rr})')
        p.number_format = MONEY
        p.font = TOT_F
        p.border = BOX
        p.alignment = CENTER
        p.fill = FILL_P2
    r1 = r0 + n - 1
    band(ws, r1 + 1, C.DEBT["totals"][0], 8, MONEY, f"=SUM(D{r0}:D{r1})")
    band(ws, r1 + 2, C.DEBT["totals"][1], 8, MONEY,
         f"=SUM(E{r0}:E{r1})+SUM(F{r0}:F{r1})")
    band(ws, r1 + 3, "TOTAL PAID OFF SO FAR", 8, MONEY,
         f"=SUM(H{r0}:H{r1})")
    note(ws, r1 + 5, C.DEBT_PLAN["quote"], 8)
    return ws


def s_week(wb, idx):
    ws = sheet(wb, f"Week {idx + 1}", [4, 13, 28, 18, 14, 14, 24])
    r = title(ws, C.WEEKLY["titles"][idx], 7)
    header_row(ws, r, C.WEEKLY["cols"], start=2)
    r0 = r + 1
    n = 22
    input_cells(ws, r0, r0 + n - 1, 2, 7)
    for rr in range(r0, r0 + n):
        ws.cell(row=rr, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=rr, column=3).alignment = LEFT
        ws.cell(row=rr, column=5).number_format = MONEY
        ws.cell(row=rr, column=7).alignment = LEFT
    r1 = r0 + n - 1
    cats = ",".join(C.WHERE["categories"])
    yesno(ws, f"D{r0}:D{r1}", f'"{cats}"')
    yesno(ws, f"F{r0}:F{r1}", '"Need,Want"')
    band(ws, r1 + 1, "TOTAL SPENT THIS WEEK", 7, MONEY, f"=SUM(E{r0}:E{r1})")
    band(ws, r1 + 2, "SPENT ON NEEDS", 7, MONEY,
         f'=SUMIF(F{r0}:F{r1},"Need",E{r0}:E{r1})')
    band(ws, r1 + 3, "SPENT ON WANTS", 7, MONEY,
         f'=SUMIF(F{r0}:F{r1},"Want",E{r0}:E{r1})')
    return ws, f"'Week {idx + 1}'!G{r1 + 1}"


def s_where(wb, week_refs):
    ws = sheet(wb, "Where It Went", [4, 24, 20, 18, 24])
    r = title(ws, C.WHERE["title"], 5,
              "Weekly totals flow in automatically from the Week tabs.")
    wr = r
    band(ws, wr, "TOTAL SPENT ACROSS WEEKS 1-4", 5, MONEY,
         "=" + "+".join(f"N({x})" for x in week_refs))
    total_cell = f"E{wr}"
    r = wr + 2
    header_row(ws, r, C.WHERE["cols"], start=2)
    r0 = r + 1
    cats = C.WHERE["categories"]
    for i, cat in enumerate(cats):
        rr = r0 + i
        c = ws.cell(row=rr, column=2, value=cat)
        c.font = LBL_F
        c.alignment = LEFT
        c.border = BOX
        c.fill = FILL_P if i % 2 else FILL_IN
        a = ws.cell(row=rr, column=3)
        a.border = BOX
        a.number_format = MONEY
        a.alignment = CENTER
        a.fill = FILL_P if i % 2 else FILL_IN
        r1p = r0 + len(cats) - 1
        p = ws.cell(row=rr, column=4,
                    value=f'=IF(SUM($C${r0}:$C${r1p})=0,"",'
                          f'C{rr}/SUM($C${r0}:$C${r1p}))')
        p.number_format = PCT
        p.font = TOT_F
        p.border = BOX
        p.alignment = CENTER
        p.fill = FILL_P2
        n = ws.cell(row=rr, column=5)
        n.border = BOX
        n.alignment = LEFT
        n.fill = FILL_P if i % 2 else FILL_IN
        ws.row_dimensions[rr].height = 21
    r1 = r0 + len(cats) - 1
    tr = r1 + 1
    for cc in range(1, 6):
        ws.cell(row=tr, column=cc).fill = FILL_P2
    lt = ws.cell(row=tr, column=2, value="TOTAL")
    lt.font = TOT_F
    lt.border = BOX
    lt.alignment = LEFT
    tv = ws.cell(row=tr, column=3, value=f"=SUM(C{r0}:C{r1})")
    tv.number_format = MONEY
    tv.font = TOT_F
    tv.border = BOX
    tv.alignment = CENTER
    tp = ws.cell(row=tr, column=4, value=f'=IF(SUM(C{r0}:C{r1})=0,"",1)')
    tp.number_format = PCT
    tp.font = TOT_F
    tp.border = BOX
    tp.alignment = CENTER
    ws.row_dimensions[tr].height = 24
    band(ws, r1 + 2, "LARGEST CATEGORY", 5, None,
         f'=IF(SUM(C{r0}:C{r1})=0,"",'
         f'INDEX(B{r0}:B{r1},MATCH(MAX(C{r0}:C{r1}),C{r0}:C{r1},0)))')
    return ws, f"'Where It Went'!C{r1 + 1}"


def s_dashboard(wb, budget_rows):
    r0, r1 = budget_rows
    ws = sheet(wb, "Dashboard", [4, 30, 22, 6, 30, 22], tab=ROSE)
    r = title(ws, C.DASHBOARD["title"], 6,
              "Everything here is calculated from the other tabs.")
    inc = f"'Monthly Budget'!D{r0}"
    tot_act = f"SUM('Monthly Budget'!D{r0}:D{r1})"
    exp = f"({tot_act}-N({inc}))"
    tiles = [
        ("TOTAL INCOME", f"=N({inc})", MONEY),
        ("TOTAL EXPENSES", f"={exp}", MONEY),
        ("TOTAL SAVINGS", f"=N('Monthly Budget'!D{r0 + 9})", MONEY),
        ("TOTAL DEBT PAYMENT", f"=N('Monthly Budget'!D{r0 + 8})", MONEY),
        ("SAVINGS RATE",
         f"=IF(N({inc})=0,\"\",N('Monthly Budget'!D{r0 + 9})/N({inc}))", PCT),
        ("MONEY LEFT OVER", f"=N({inc})-{exp}", MONEY),
    ]
    rr = r
    for i, (lbl, formula, fmt) in enumerate(tiles):
        col = 2 if i % 2 == 0 else 5
        row = rr + (i // 2) * 3
        L = ws.cell(row=row, column=col, value=lbl)
        L.font = SEC_F
        L.fill = FILL_P
        L.alignment = LEFT
        L.border = BOX
        v = ws.cell(row=row, column=col + 1, value=formula)
        v.number_format = fmt
        v.font = Font(name="Calibri", size=13, bold=True, color=BROWN)
        v.fill = FILL_P
        v.border = BOX
        v.alignment = CENTER
        ws.row_dimensions[row].height = 28
    rr = rr + 9
    for lbl in ["BIGGEST FINANCIAL WIN", "MAIN FOCUS FOR NEXT MONTH"]:
        ws.cell(row=rr, column=2, value=lbl).font = SEC_F
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
        c = ws.cell(row=rr, column=3)
        c.fill = FILL_IN
        c.border = BOX
        c.alignment = LEFT
        ws.row_dimensions[rr].height = 26
        rr += 2
    rr += 1
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
    ws.cell(row=rr, column=2, value=C.FINAL["statement"]).font = Font(
        name="Georgia", size=16, bold=True, color=BROWN)
    ws.cell(row=rr, column=2).alignment = CENTER
    rr += 2
    for p in C.REFLECTION["prompts"]:
        ws.cell(row=rr, column=2, value=p).font = LBL_F
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
        c = ws.cell(row=rr, column=3)
        c.fill = FILL_IN
        c.border = BOX
        c.alignment = LEFT
        ws.row_dimensions[rr].height = 24
        rr += 1
    return ws


def s_reset(wb):
    ws = sheet(wb, "Monthly Reset", [4, 8, 42, 22])
    r = title(ws, C.RESET["title"], 4)
    header_row(ws, r, ["", "TASK", "DONE"], start=2)
    r0 = r + 1
    for i, it in enumerate(C.RESET["items"]):
        rr = r0 + i
        n = ws.cell(row=rr, column=2, value=i + 1)
        n.font = LBL_F
        n.alignment = CENTER
        n.border = BOX
        c = ws.cell(row=rr, column=3, value=it)
        c.font = BODY_F
        c.alignment = LEFT
        c.border = BOX
        d = ws.cell(row=rr, column=4)
        d.border = BOX
        d.alignment = CENTER
        for cc in (2, 3, 4):
            ws.cell(row=rr, column=cc).fill = FILL_P if i % 2 else FILL_IN
        ws.row_dimensions[rr].height = 23
    r1 = r0 + len(C.RESET["items"]) - 1
    yesno(ws, f"D{r0}:D{r1}", '"Done,Not yet"')
    band(ws, r1 + 1, "COMPLETED", 4, "0",
         f'=COUNTIF(D{r0}:D{r1},"Done")&" of {len(C.RESET["items"])}"')
    note(ws, r1 + 3, C.RESET["quote"], 4)

    rr = r1 + 5
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=4)
    ws.cell(row=rr, column=2, value=C.NEXT_MONTH["title"]).font = Font(
        name="Georgia", size=14, bold=True, color=BROWN)
    ws.cell(row=rr, column=2).alignment = CENTER
    rr += 1
    for s in C.NEXT_MONTH["sections"]:
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        c = ws.cell(row=rr, column=2, value=s)
        c.font = LBL_F
        c.alignment = LEFT
        v = ws.cell(row=rr, column=4)
        v.fill = FILL_IN
        v.border = BOX
        ws.row_dimensions[rr].height = 24
        rr += 1
    return ws


def main():
    os.makedirs(OUT, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    s_start(wb)
    s_year(wb)
    s_income(wb)
    s_fixed(wb)
    s_variable(wb)
    _, brows = None, None
    ws_b, br0, br1 = s_budget(wb, None)
    s_bills(wb)
    s_subs(wb)
    s_savings(wb)
    s_challenge(wb)
    s_debt(wb)
    week_refs = []
    for i in range(4):
        _, ref = s_week(wb, i)
        week_refs.append(ref)
    s_where(wb, week_refs)
    s_dashboard(wb, (br0, br1))
    s_reset(wb)

    wb.properties.title = ("Monthly Budget Planner \u2014 Editable Workbook")
    wb.properties.creator = "Cedric Errol Cajelo"
    wb.properties.subject = "Editable monthly budget planner"

    path = os.path.join(OUT, "budget_planner_EDITABLE.xlsx")
    wb.save(path)
    print(f"{os.path.basename(path)}: {len(wb.sheetnames)} sheets  "
          f"{os.path.getsize(path)/1024:.0f} KB")
    print("  " + " | ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
