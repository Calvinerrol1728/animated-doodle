#!/usr/bin/env python3
"""
Preflight for the MONTHLY BUDGET PLANNER.

Checks the finished files against the brief rather than against my own
intentions: page count, trim sizes, embedded fonts, image resolution, real
vector text, verbatim wording, no page numbers, no stray placeholder text,
working form fields, and the absence of any invented financial figure.
"""
import os
import re
import sys
import unicodedata

import pymupdf
from pypdf import PdfReader

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")

LETTER = os.path.join(OUT, "budget_planner_US_Letter.pdf")
A4 = os.path.join(OUT, "budget_planner_A4.pdf")
FILL = os.path.join(OUT, "budget_planner_FILLABLE.pdf")
FILL_A4 = os.path.join(OUT, "budget_planner_FILLABLE_A4.pdf")
XLSX = os.path.join(OUT, "budget_planner_EDITABLE.xlsx")

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}"
          + (f"  \u2014 {detail}" if detail else ""))
    return ok


def norm(s):
    s = unicodedata.normalize("NFKD", s)
    for a, b in [("\u2019", "'"), ("\u2018", "'"), ("\u201c", '"'),
                 ("\u201d", '"'), ("\u2014", "-"), ("\u2013", "-"),
                 ("\u2022", "*")]:
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip().lower()


def all_strings():
    """Every phrase the brief says must appear, flattened."""
    out = []

    def walk(v):
        if isinstance(v, str):
            out.append(v)
        elif isinstance(v, dict):
            for x in v.values():
                walk(x)
        elif isinstance(v, (list, tuple)):
            for x in v:
                walk(x)

    for key in dir(C):
        if key.isupper():
            walk(getattr(C, key))
    return [s for s in out if len(s.strip()) > 2]


def main():
    print("\nMONTHLY BUDGET PLANNER \u2014 PREFLIGHT\n" + "=" * 70)

    for path in (LETTER, A4, FILL, FILL_A4, XLSX):
        check(f"file exists: {os.path.basename(path)}", os.path.exists(path),
              f"{os.path.getsize(path)/1e6:.2f} MB" if os.path.exists(path)
              else "MISSING")

    print("\nPRINT EDITIONS")
    for path, w_in, h_in in ((LETTER, 8.5, 11.0), (A4, 8.268, 11.693)):
        r = PdfReader(path)
        base = os.path.basename(path)
        check(f"{base}: 40 pages", len(r.pages) == 40, f"{len(r.pages)}")
        sizes = {(round(float(p.mediabox.width) / 72, 3),
                  round(float(p.mediabox.height) / 72, 3)) for p in r.pages}
        check(f"{base}: uniform trim size", len(sizes) == 1, str(sizes))
        got = sizes.pop()
        check(f"{base}: correct trim",
              abs(got[0] - w_in) < 0.02 and abs(got[1] - h_in) < 0.02,
              f"{got[0]} x {got[1]} in")

        # fonts embedded
        d = pymupdf.open(path)
        unembedded = set()
        for pno in range(d.page_count):
            for f in d.get_page_fonts(pno):
                if not f[1]:
                    unembedded.add(f[3])
        check(f"{base}: all fonts embedded", not unembedded, str(unembedded))

        # image resolution
        worst = None
        for pno in range(d.page_count):
            for img in d[pno].get_images(full=True):
                xref = img[0]
                info = d.extract_image(xref)
                for rect in d[pno].get_image_rects(xref):
                    dpi = min(info["width"] / (rect.width / 72),
                              info["height"] / (rect.height / 72))
                    worst = dpi if worst is None else min(worst, dpi)
        check(f"{base}: art at 300 DPI or better",
              worst is None or worst >= 295,
              "no raster art" if worst is None else f"min {worst:.0f} DPI")

        # real vector text
        chars = sum(len(d[p].get_text().strip()) for p in range(d.page_count))
        check(f"{base}: text is live vector type", chars > 4000,
              f"{chars} characters")

        # no page numbers: a page whose only bottom content is a bare number
        offenders = []
        for pno in range(d.page_count):
            page = d[pno]
            bottom = page.rect.height - 40
            for blk in page.get_text("blocks"):
                if blk[1] > bottom and re.fullmatch(r"\s*\d{1,3}\s*", blk[4]):
                    offenders.append(pno + 1)
        check(f"{base}: no page numbers", not offenders, str(offenders))

        # placeholder / lorem / TODO
        full = " ".join(d[p].get_text() for p in range(d.page_count))
        junk = [w for w in ("lorem", "ipsum", "todo", "tbd", "xxx",
                            "placeholder", "sample text")
                if w in full.lower()]
        check(f"{base}: no placeholder text", not junk, str(junk))

        # no invented money figures: any currency amount is a defect
        money = re.findall(r"[$€£]\s?\d", full)
        check(f"{base}: no invented figures", not money, str(money[:5]))

        # no advice/guarantee language
        banned = ["guaranteed", "guarantee", "risk-free", "get rich",
                  "double your money", "invest in", "returns of",
                  "financial advice", "we recommend investing"]
        hits = [b for b in banned if b in full.lower()]
        check(f"{base}: no advice or guarantees", not hits, str(hits))
        d.close()

    print("\nVERBATIM TEXT")
    d = pymupdf.open(LETTER)
    body = norm(" ".join(d[p].get_text() for p in range(d.page_count)))
    d.close()
    strings = all_strings()
    missing = [s for s in strings if norm(s) not in body]
    check(f"all {len(strings)} required phrases present verbatim",
          not missing, f"{len(missing)} missing: {missing[:4]}")

    print("\nREQUIRED SECTIONS")
    must = ["MONTHLY BUDGET PLANNER", "THIS PLANNER BELONGS TO",
            "HOW TO USE YOUR PLANNER", "MY FINANCIAL VISION",
            "MY YEAR AT A GLANCE", "INCOME TRACKER", "FIXED EXPENSES",
            "VARIABLE EXPENSES", "MY MONTHLY BUDGET", "NEEDS VS. WANTS",
            "BILL PAYMENT TRACKER", "SUBSCRIPTION CHECKUP",
            "MY SAVINGS GOALS", "EMERGENCY FUND", "DEBT TRACKER",
            "WEEK 1", "WEEK 2", "WEEK 3", "WEEK 4", "NO-SPEND CHALLENGE",
            "GROCERY & MEAL BUDGET", "SHOPPING PLANNER", "MY MONEY HABITS",
            "WEEKLY MONEY CHECK-IN", "MONTH-END BUDGET REVIEW",
            "WHERE DID MY MONEY GO?", "MY SAVINGS REVIEW",
            "CELEBRATE YOUR FINANCIAL WINS", "MONTHLY MONEY REFLECTION",
            "MY FINANCIAL DASHBOARD", "MY MONEY GOALS",
            "MY FINANCIAL RESET", "MY MONEY HAS A PLAN"]
    absent = [m for m in must if norm(m) not in body]
    check(f"all {len(must)} named sections present", not absent, str(absent))

    print("\nFILLABLE EDITIONS")
    for path in (FILL, FILL_A4):
        d = pymupdf.open(path)
        base = os.path.basename(path)
        check(f"{base}: 40 pages", d.page_count == 40, str(d.page_count))
        n = sum(len(list(d[p].widgets())) for p in range(d.page_count))
        check(f"{base}: form fields present", n > 1000, f"{n} fields")
        pages_with = sum(1 for p in range(d.page_count)
                         if list(d[p].widgets()))
        check(f"{base}: every writing page is fillable", pages_with >= 37,
              f"{pages_with}/40 pages")
        toc = d.get_toc()
        check(f"{base}: bookmark navigation", len(toc) >= 40, f"{len(toc)}")
        # overlap
        bad = 0
        for p in range(d.page_count):
            rs = [w.rect for w in d[p].widgets()]
            for i in range(len(rs)):
                for j in range(i + 1, len(rs)):
                    it = rs[i] & rs[j]
                    if it.is_valid and it.get_area() > 4:
                        bad += 1
        check(f"{base}: no overlapping fields", bad == 0, f"{bad} pairs")
        d.close()

    print("\nEDITABLE WORKBOOK")
    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    check("workbook opens", True, f"{len(wb.sheetnames)} sheets")
    check("has 18 tabs", len(wb.sheetnames) == 18, str(len(wb.sheetnames)))
    formulas = 0
    hardcoded = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                elif isinstance(c.value, (int, float)):
                    # only currency-formatted cells could be a fake figure;
                    # plain integers are day and row numbers
                    if "$" in (c.number_format or ""):
                        hardcoded.append(
                            f"{ws.title}!{c.coordinate}={c.value}")
    check("live formulas throughout", formulas >= 120,
          f"{formulas} formulas")
    check("no pre-filled money values", not hardcoded, str(hardcoded[:5]))
    names = ["Start Here", "Dashboard", "Monthly Budget", "Where It Went"]
    check("key tabs present",
          all(n in wb.sheetnames for n in names), str(wb.sheetnames))

    print("\n" + "=" * 70)
    ok = sum(1 for _n, o, _d in results if o)
    total = len(results)
    print(f"{ok}/{total} checks passed"
          + ("  \u2014 ALL CLEAR" if ok == total else "  \u2014 SEE FAILURES ABOVE"))
    return total - ok


if __name__ == "__main__":
    raise SystemExit(1 if main() else 0)
