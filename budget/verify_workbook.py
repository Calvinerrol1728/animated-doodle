#!/usr/bin/env python3
"""
Verification harness for the editable workbook.

Fills every sheet with sample figures, recalculates the whole file with a real
formula engine, and checks each computed cell against an independently
calculated expected value. Nothing is trusted just because it looks right.
"""
import logging
import os
import shutil
import warnings

warnings.filterwarnings("ignore")
logging.disable(logging.CRITICAL)

import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "out", "budget_planner_EDITABLE.xlsx")
TMP = "/tmp/verify_wb.xlsx"


def find_row(ws, text, col=2, limit=80):
    t = str(text).strip().lower()
    for r in range(1, limit):
        v = ws.cell(row=r, column=col).value
        if v and str(v).strip().lower().startswith(t):
            return r
    return None


def find_band(ws, text, limit=80):
    """Merged total bands are written into column 1."""
    t = str(text).strip().lower()
    for r in range(1, limit):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().lower() == t:
            return r
    return None


def main():
    shutil.copy(SRC, TMP)
    wb = openpyxl.load_workbook(TMP)
    exp = {}

    # ---------------- Year at a Glance
    ws = wb["Year at a Glance"]
    r0 = find_row(ws, "January")
    for i in range(12):
        ws.cell(row=r0 + i, column=3, value=4000)
        ws.cell(row=r0 + i, column=4, value=3000)
        ws.cell(row=r0 + i, column=5, value=500)
    tr = find_row(ws, "TOTAL")
    exp[("Year at a Glance", f"F{r0}")] = 500
    exp[("Year at a Glance", f"C{tr}")] = 48000
    exp[("Year at a Glance", f"F{tr}")] = 6000

    # ---------------- Income
    ws = wb["Income"]
    r0 = find_row(ws, "TOTAL EXPECTED INCOME", col=1)
    hdr = None
    for r in range(1, 40):
        if ws.cell(row=r, column=2).value == "DATE":
            hdr = r
            break
    d0 = hdr + 1
    for i in range(5):
        ws.cell(row=d0 + i, column=4, value=1000)
        ws.cell(row=d0 + i, column=5, value=950)
    b1 = find_band(ws, "TOTAL EXPECTED INCOME:")
    b2 = find_band(ws, "TOTAL ACTUAL INCOME:")
    b3 = find_band(ws, "DIFFERENCE")
    exp[("Income", f"F{b1}")] = 5000
    exp[("Income", f"F{b2}")] = 4750
    exp[("Income", f"F{b3}")] = -250

    # ---------------- Fixed Expenses
    ws = wb["Fixed Expenses"]
    hdr = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "EXPENSE")
    d0 = hdr + 1
    for i in range(4):
        ws.cell(row=d0 + i, column=4, value=200)
        ws.cell(row=d0 + i, column=5, value=210)
    exp[("Fixed Expenses", f"F{find_band(ws, 'TOTAL BUDGETED')}")] = 800
    exp[("Fixed Expenses", f"F{find_band(ws, 'TOTAL ACTUAL')}")] = 840

    # ---------------- Variable Expenses
    ws = wb["Variable Expenses"]
    r0 = find_row(ws, "FOOD")
    for i in range(6):
        ws.cell(row=r0 + i, column=3, value=300)
        ws.cell(row=r0 + i, column=4, value=250)
    exp[("Variable Expenses", f"E{r0}")] = 50
    exp[("Variable Expenses", f"E{find_band(ws, 'TOTAL')}")] = 1500

    # ---------------- Monthly Budget
    ws = wb["Monthly Budget"]
    vals = {"INCOME": 4200, "HOUSING": 1500, "UTILITIES": 220, "FOOD": 600,
            "TRANSPORTATION": 180, "HEALTH": 120, "PERSONAL": 90,
            "ENTERTAINMENT": 110, "DEBT": 300, "SAVINGS": 500, "OTHER": 60}
    rows = {}
    for name, v in vals.items():
        r = find_row(ws, name)
        rows[name] = r
        ws.cell(row=r, column=3, value=v + 100)
        ws.cell(row=r, column=4, value=v)
    tr = find_row(ws, "TOTAL")
    exp[("Monthly Budget", f"E{rows['HOUSING']}")] = 100
    exp[("Monthly Budget", f"D{tr}")] = sum(vals.values())
    exp[("Monthly Budget", f"E{tr}")] = 1100

    # ---------------- Bills
    ws = wb["Bills"]
    hdr = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "BILL")
    d0 = hdr + 1
    for i in range(6):
        ws.cell(row=d0 + i, column=4, value=100)
        ws.cell(row=d0 + i, column=6, value="Paid" if i < 4 else "Due")
    exp[("Bills", f"F{find_band(ws, 'TOTAL BILLS')}")] = 600
    exp[("Bills", f"F{find_band(ws, 'NUMBER MARKED PAID')}")] = 4

    # ---------------- Subscriptions
    ws = wb["Subscriptions"]
    hdr = next(r for r in range(1, 40)
               if ws.cell(row=r, column=2).value == "SUBSCRIPTION")
    d0 = hdr + 1
    for i in range(6):
        ws.cell(row=d0 + i, column=3, value=10)
        ws.cell(row=d0 + i, column=5, value="Keep" if i % 2 == 0 else "Cancel")
    exp[("Subscriptions", f"F{find_band(ws, 'TOTAL PER MONTH')}")] = 60
    exp[("Subscriptions", f"F{find_band(ws, 'TOTAL PER YEAR')}")] = 720
    exp[("Subscriptions",
         f"F{find_band(ws, 'MONTHLY COST OF THE ONES MARKED KEEP')}")] = 30
    exp[("Subscriptions",
         f"F{find_band(ws, 'MONTHLY SAVING IF YOU CANCEL THE REST')}")] = 30

    # ---------------- Savings Goals + Emergency fund
    ws = wb["Savings Goals"]
    hdr = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "GOAL")
    d0 = hdr + 1
    ws.cell(row=d0, column=3, value=2000)
    ws.cell(row=d0, column=4, value=750)
    ws.cell(row=d0 + 1, column=3, value=1000)
    ws.cell(row=d0 + 1, column=4, value=250)
    exp[("Savings Goals", f"F{d0}")] = 0.375
    exp[("Savings Goals", f"F{find_band(ws, 'TOTAL SAVED TOWARD GOALS')}")] = 1000
    exp[("Savings Goals", f"F{find_band(ws, 'TOTAL STILL TO GO')}")] = 2000
    er = find_row(ws, "My current emergency fund")
    ws.cell(row=er, column=4, value=1200)
    ws.cell(row=er + 1, column=4, value=6000)
    exp[("Savings Goals", f"D{er + 3}")] = 0.2

    # ---------------- Savings Challenge
    ws = wb["Savings Challenge"]
    hdr = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "DAY")
    d0 = hdr + 1
    tgt = find_row(ws, C_target := "MY 30-DAY TARGET")
    ws.cell(row=tgt, column=3, value=500)
    for i in range(15):
        ws.cell(row=d0 + i, column=3, value=5)
        ws.cell(row=d0 + i, column=7, value=10)
    exp[("Savings Challenge", f"D{d0 + 14}")] = 75          # day 15 running
    exp[("Savings Challenge", f"H{d0 + 14}")] = 225         # day 30 running
    exp[("Savings Challenge",
         f"H{find_band(ws, 'TOTAL SAVED OVER 30 DAYS')}")] = 225
    exp[("Savings Challenge",
         f"H{find_band(ws, 'AMOUNT LEFT TO REACH YOUR TARGET')}")] = 275

    # ---------------- Debt
    ws = wb["Debt"]
    hdr = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "DEBT")
    d0 = hdr + 1
    ws.cell(row=d0, column=3, value=5000)
    ws.cell(row=d0, column=4, value=4200)
    ws.cell(row=d0, column=5, value=150)
    ws.cell(row=d0, column=6, value=100)
    exp[("Debt", f"H{d0}")] = 800
    exp[("Debt", f"H{find_band(ws, 'TOTAL DEBT:')}")] = 4200
    exp[("Debt", f"H{find_band(ws, 'TOTAL PAID THIS MONTH:')}")] = 250
    exp[("Debt", f"H{find_band(ws, 'TOTAL PAID OFF SO FAR')}")] = 800

    # ---------------- Weeks
    for wi in range(1, 5):
        ws = wb[f"Week {wi}"]
        hdr = next(r for r in range(1, 40)
                   if ws.cell(row=r, column=2).value == "DATE")
        d0 = hdr + 1
        for i in range(6):
            ws.cell(row=d0 + i, column=5, value=25)
            ws.cell(row=d0 + i, column=6, value="Need" if i < 4 else "Want")
        exp[(f"Week {wi}", f"G{find_band(ws, 'TOTAL SPENT THIS WEEK')}")] = 150
        exp[(f"Week {wi}", f"G{find_band(ws, 'SPENT ON NEEDS')}")] = 100
        exp[(f"Week {wi}", f"G{find_band(ws, 'SPENT ON WANTS')}")] = 50

    # ---------------- Where It Went
    ws = wb["Where It Went"]
    hdr = next(r for r in range(1, 40)
               if ws.cell(row=r, column=2).value == "CATEGORY")
    d0 = hdr + 1
    for i in range(9):
        ws.cell(row=d0 + i, column=3, value=100 if i else 300)
    total = 300 + 8 * 100
    exp[("Where It Went", f"D{d0}")] = 300 / total
    exp[("Where It Went", f"C{d0 + 9}")] = total
    exp[("Where It Went",
         f"E{find_band(ws, 'TOTAL SPENT ACROSS WEEKS 1-4')}")] = 600
    exp[("Where It Went", f"E{find_band(ws, 'LARGEST CATEGORY')}")] = "Housing"

    # ---------------- Monthly Reset
    ws = wb["Monthly Reset"]
    d0 = None
    for r in range(1, 40):
        if ws.cell(row=r, column=3).value == "Review my income":
            d0 = r
            break
    for i in range(5):
        ws.cell(row=d0 + i, column=4, value="Done")
    exp[("Monthly Reset", f"D{find_band(ws, 'COMPLETED')}")] = "5 of 10"

    # ---------------- Dashboard expectations
    dash = wb["Dashboard"]
    exp[("Dashboard", "C4")] = 4200                     # total income
    exp[("Dashboard", "F4")] = sum(vals.values()) - 4200  # total expenses
    exp[("Dashboard", "C7")] = 500                      # savings
    exp[("Dashboard", "F7")] = 300                      # debt payment
    exp[("Dashboard", "C10")] = 500 / 4200              # savings rate
    exp[("Dashboard", "F10")] = 4200 - (sum(vals.values()) - 4200)

    wb.save(TMP)

    # ---------------- recalculate
    import formulas
    sol = formulas.ExcelModel().loads(TMP).finish().calculate()
    got = {}
    for k, v in sol.items():
        if "'!" not in k:
            continue
        sheet = k.split("]")[-1].split("'!")[0].strip("'")
        cell = k.split("'!")[-1].rstrip("]")
        if ":" in cell:
            continue
        try:
            got[(sheet.upper(), cell)] = v.value[0, 0]
        except Exception:
            pass

    ok = bad = 0
    print(f"{'CHECK':<58}{'EXPECTED':>14}{'GOT':>16}   ")
    print("-" * 92)
    for (sheet, cell), want in sorted(exp.items()):
        have = got.get((sheet.upper(), cell), "<missing>")
        if isinstance(want, float) and isinstance(have, (int, float)):
            good = abs(float(have) - want) < 1e-6
        elif isinstance(want, (int,)) and isinstance(have, (int, float)):
            good = abs(float(have) - want) < 1e-6
        else:
            good = str(have).strip() == str(want).strip()
        mark = "PASS" if good else "FAIL"
        ok, bad = (ok + 1, bad) if good else (ok, bad + 1)
        wd = f"{want:.4f}" if isinstance(want, float) else str(want)
        hd = f"{have:.4f}" if isinstance(have, float) else str(have)
        print(f"{sheet + ' ' + cell:<58}{wd:>14}{hd:>16}   {mark}")
    print("-" * 92)
    print(f"{ok}/{ok + bad} formula checks passed"
          + ("" if bad == 0 else f"  \u2014 {bad} FAILED"))
    return bad


if __name__ == "__main__":
    raise SystemExit(1 if main() else 0)
